"""
Local browser-based Python code runner.

Run with:
    python app.py

Then open http://localhost:5000 in your browser.

SECURITY NOTE:
This runs whatever Python code is submitted, on your own machine,
with your own user permissions. It is intended for LOCAL, PERSONAL
use only (e.g. testing snippets on your own laptop). Do NOT expose
this server to the network or the internet — anyone who can reach
it can run arbitrary code on your computer.
"""

import subprocess
import sys
import tempfile
import os
import uuid
import json
import shlex
import time
import threading
import logging
import logging.handlers
import re
import html
import math
import base64
from io import BytesIO
from datetime import datetime, timezone, timedelta
import xml.etree.ElementTree as ET
from xml.dom import minidom

import requests
from flask import Flask, request, jsonify, render_template, send_from_directory, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# Where compiled executables, config, and logs are saved. Everything
# lives next to this file.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
BUILDS_DIR = os.path.join(BASE_DIR, "builds")
os.makedirs(BUILDS_DIR, exist_ok=True)

# Log file path — only written to while logging is enabled (see
# configure_logging below). Rotated so it can't grow unbounded over a
# long-running session.
LOG_FILE_PATH = os.path.join(BASE_DIR, "app.log")

log = logging.getLogger("app")


def configure_logging(enabled):
    """Turns logging to the terminal and to app.log on or off. Affects
    the root logger, so it covers both our own log.info() calls (e.g.
    diagnosing third-party API calls to Airlock Digital, VirusTotal,
    the Cloud multi-tenant API) and Flask/Werkzeug's own request
    logging — a single toggle for everything currently visible in the
    terminal. Safe to call again at runtime to flip the setting
    without restarting the app."""
    root_logger = logging.getLogger()
    for handler in list(root_logger.handlers):
        root_logger.removeHandler(handler)

    if not enabled:
        # A NullHandler avoids Python's "no handlers found" warning;
        # setting the level above CRITICAL means log calls are cheap
        # no-ops rather than doing real work that goes nowhere.
        root_logger.addHandler(logging.NullHandler())
        root_logger.setLevel(logging.CRITICAL + 1)
        return

    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s")

    console_handler = logging.StreamHandler()
    console_handler.setFormatter(formatter)
    root_logger.addHandler(console_handler)

    file_handler = logging.handlers.RotatingFileHandler(
        LOG_FILE_PATH, maxBytes=5 * 1024 * 1024, backupCount=3, encoding="utf-8"
    )
    file_handler.setFormatter(formatter)
    root_logger.addHandler(file_handler)

    root_logger.setLevel(logging.INFO)


# On by default (matches this app's behavior before the toggle existed)
# so logging works immediately for any use of this module, including
# being imported directly for testing. __main__ below re-applies the
# actual saved preference once config.json can be read.
configure_logging(True)

# Max time (seconds) a submitted script is allowed to run before being killed.
EXECUTION_TIMEOUT = 10

# Max time (seconds) a compile job is allowed to run before being killed.
# PyInstaller can be slow, especially on the first run.
COMPILE_TIMEOUT = 180

# Local config file storing the VirusTotal API key.
# NOTE: stored in plaintext. Fine for a single-user local tool, but
# don't commit this file or share it — it grants access to your
# VirusTotal account/quota.
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")

# Persisted Timed Audit Mode sessions, so scheduled reverts survive an
# app restart. Local file next to app.py; not sensitive on its own
# (agent IDs and group IDs only, no secrets) but still git-ignored.
AUDIT_SESSIONS_PATH = os.path.join(BASE_DIR, "audit_sessions.json")

# Editable ISO 27001 <-> Airlock policy control mapping. This is
# reference content, not a secret or runtime state — it's meant to be
# committed to git and hand-tuned per client engagement.
ISO_MAPPING_PATH = os.path.join(BASE_DIR, "iso_mapping.json")

# NFR Tracking widget's own storage — deliberately separate from
# config.json so its data can never collide with anything else the app
# writes. Contains the Cloud admin credential and saved partner
# tenants.
CLOUD_CONFIG_PATH = os.path.join(BASE_DIR, "cloud_config.json")

VT_FILES_URL = "https://www.virustotal.com/api/v3/files"
VT_UPLOAD_URL_ENDPOINT = "https://www.virustotal.com/api/v3/files/upload_url"
VT_LARGE_FILE_THRESHOLD = 32 * 1024 * 1024  # VT requires the special upload URL above 32MB

# Max time (seconds) an API-calls script is allowed to run before being killed.
DEFAULT_API_SCRIPT_TIMEOUT = 120

# Default folder for your Airlock Digital script repository, used when
# api_scripts_dir isn't set in config.json. You can point this at your
# existing repo instead via the Settings tab. Scripts always run from
# this local folder — GitHub, when enabled, is a sync source that copies
# .py files into this folder, not a separate execution location.
DEFAULT_API_SCRIPTS_DIR = os.path.join(BASE_DIR, "api_scripts")

# File types the app will offer to open after an API-calls script runs.
OUTPUT_FILE_EXTENSIONS = (".xlsx", ".xml", ".html", ".htm")

# Labels scripts commonly use to announce where they wrote their output,
# most specific first. Matched case-insensitively against each line of
# stdout/stderr, e.g. "HTML report: report\policy_report.html".
OUTPUT_LABEL_PATTERNS = [
    re.compile(r"html report\s*:\s*(.+)", re.IGNORECASE),
    re.compile(r"report file\s*:\s*(.+)", re.IGNORECASE),
    re.compile(r"output file\s*:\s*(.+)", re.IGNORECASE),
    re.compile(r"\bsaved\s*:\s*(.+)", re.IGNORECASE),
    re.compile(r"\breport\s*:\s*(.+)", re.IGNORECASE),
    re.compile(r"\boutput\s*:\s*(.+)", re.IGNORECASE),
]


def extract_output_path_from_text(text, exec_dir):
    """Looks for a line like 'HTML report: some\\path\\file.html' in a
    script's console output and resolves it to a real file under
    exec_dir. Returns the absolute path, or None if nothing matched or
    the resulting file doesn't actually exist.

    Checks label patterns in priority order (most specific label first)
    across the *entire* output before falling back to a less-specific
    label — so a specific "HTML report:" line always wins over a generic
    "Saved:" line from an earlier step (e.g. an intermediate XML export),
    even though that generic line appears first in the text.
    """
    if not text:
        return None

    real_exec_dir = os.path.realpath(exec_dir)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

    def resolve(candidate):
        candidate = candidate.strip().strip('"').strip("'")
        if not candidate.lower().endswith(OUTPUT_FILE_EXTENSIONS):
            return None

        # Scripts run on Windows commonly print backslash paths.
        normalized = candidate.replace("\\", os.sep).replace("/", os.sep)
        full_path = normalized if os.path.isabs(normalized) else os.path.join(exec_dir, normalized)
        full_path = os.path.normpath(full_path)

        # Security: only accept paths that stay inside the scripts folder.
        real_full = os.path.realpath(full_path)
        if real_full != real_exec_dir and not real_full.startswith(real_exec_dir + os.sep):
            return None

        return full_path if os.path.isfile(full_path) else None

    for pattern in OUTPUT_LABEL_PATTERNS:
        # If a pattern matches more than one line (e.g. a script logs its
        # own "Saved:" line for each step), prefer the last one — final
        # summary lines are typically printed after the interim ones.
        match = None
        for line in lines:
            m = pattern.search(line)
            if not m:
                continue
            resolved = resolve(m.group(1))
            if resolved:
                match = resolved
        if match:
            return match

    return None


def load_config():
    if os.path.exists(CONFIG_PATH):
        try:
            with open(CONFIG_PATH, "r") as f:
                return json.load(f)
        except (json.JSONDecodeError, OSError):
            return {}
    return {}


def save_config(config):
    with open(CONFIG_PATH, "w") as f:
        json.dump(config, f, indent=2)


def load_cloud_config():
    if os.path.exists(CLOUD_CONFIG_PATH):
        try:
            with open(CLOUD_CONFIG_PATH, "r") as f:
                return json.load(f)
        except (json.JSONDecodeError, OSError):
            return {}
    return {}


def save_cloud_config(cfg):
    with open(CLOUD_CONFIG_PATH, "w") as f:
        json.dump(cfg, f, indent=2)
    log.info("Wrote %s (%d tenant(s))", CLOUD_CONFIG_PATH, len((cfg.get("cloud_tenants") or [])))


def migrate_legacy_cloud_config():
    """One-time migration: cloud_admin/cloud_tenants used to live inside
    the main config.json. Move them into their own dedicated file so
    this widget's data is fully isolated from everything else the app
    writes — never touched by, or touching, unrelated saves."""
    config = load_config()
    if "cloud_admin" not in config and "cloud_tenants" not in config:
        return  # nothing to migrate

    cloud_cfg = load_cloud_config()
    if "cloud_admin" in config and not cloud_cfg.get("cloud_admin"):
        cloud_cfg["cloud_admin"] = config["cloud_admin"]
    if "cloud_tenants" in config and not cloud_cfg.get("cloud_tenants"):
        cloud_cfg["cloud_tenants"] = config["cloud_tenants"]
    save_cloud_config(cloud_cfg)

    config.pop("cloud_admin", None)
    config.pop("cloud_tenants", None)
    save_config(config)
    print(f"[startup] Migrated NFR Tracking settings into their own file: {CLOUD_CONFIG_PATH}")


# Guards audit_sessions.json against concurrent access between HTTP
# request handlers and the background revert-scheduler thread.
AUDIT_SESSIONS_LOCK = threading.Lock()


def load_audit_sessions():
    with AUDIT_SESSIONS_LOCK:
        if os.path.exists(AUDIT_SESSIONS_PATH):
            try:
                with open(AUDIT_SESSIONS_PATH, "r") as f:
                    return json.load(f)
            except (json.JSONDecodeError, OSError):
                return []
        return []


def save_audit_sessions(sessions):
    with AUDIT_SESSIONS_LOCK:
        with open(AUDIT_SESSIONS_PATH, "w") as f:
            json.dump(sessions, f)


def get_scripts_dir():
    configured = load_config().get("api_scripts_dir", "").strip()
    scripts_dir = configured if configured else DEFAULT_API_SCRIPTS_DIR
    os.makedirs(scripts_dir, exist_ok=True)
    return scripts_dir


def get_script_timeout():
    configured = load_config().get("api_script_timeout", "").strip()
    if configured.isdigit() and int(configured) > 0:
        return int(configured)
    return DEFAULT_API_SCRIPT_TIMEOUT


# Default port for the Airlock Digital REST API, per Airlock's own docs.
DEFAULT_AIRLOCK_PORT = 3129


def get_active_airlock_profile(config=None):
    """Returns the currently active saved Airlock connection (a dict
    with id/label/tenant/port/api_key), or None if none is set."""
    config = config if config is not None else load_config()
    profiles = config.get("airlock_profiles") or []
    active_id = config.get("active_airlock_profile_id")
    for p in profiles:
        if p.get("id") == active_id:
            return p
    return None


def migrate_legacy_airlock_profile():
    """One-time migration: older versions of this app stored a single
    flat airlock_tenant/airlock_port/airlock_api_key in config.json.
    If that's all that's there and no connections have been saved yet,
    convert it into the first saved connection so nobody loses their
    existing setup on upgrade."""
    config = load_config()
    if config.get("airlock_profiles"):
        return  # already migrated (or already using the new system)

    tenant = (config.get("airlock_tenant") or "").strip()
    api_key = (config.get("airlock_api_key") or "").strip()
    if not tenant and not api_key:
        return  # nothing to migrate

    profile = {
        "id": str(uuid.uuid4()),
        "label": "Default",
        "tenant": tenant,
        "port": (config.get("airlock_port") or "").strip(),
        "api_key": api_key,
    }
    config["airlock_profiles"] = [profile]
    config["active_airlock_profile_id"] = profile["id"]
    config.pop("airlock_tenant", None)
    config.pop("airlock_api_key", None)
    config.pop("airlock_port", None)
    save_config(config)
    print(f"[startup] Migrated existing Airlock Digital settings into a saved connection: '{profile['label']}'")


def get_airlock_base_url():
    """Returns (base_url, error) based on the active saved connection."""
    profile = get_active_airlock_profile()
    if not profile:
        return None, "No active Airlock Digital connection. Add or select one in Settings."
    tenant = (profile.get("tenant") or "").strip()
    if not tenant:
        return None, "The active Airlock connection has no tenant set."
    port = (profile.get("port") or "").strip() or str(DEFAULT_AIRLOCK_PORT)
    # Allow the tenant field to already include a port (tenant:port) without
    # doubling up, in case someone pastes it in that form.
    host = tenant.split(":")[0]
    return f"https://{host}:{port}", None


def airlock_request(path, payload=None):
    """POSTs to the Airlock Digital REST API and returns (data, error).
    data is the parsed 'response' object on success; error is a plain,
    display-ready message on any failure (config, network, auth, or API-
    level error)."""
    base_url, err = get_airlock_base_url()
    if err:
        return None, err

    profile = get_active_airlock_profile()
    api_key = (profile.get("api_key") or "").strip() if profile else ""
    if not api_key:
        return None, "The active Airlock connection has no API key set."

    try:
        resp = requests.post(
            f"{base_url}{path}",
            json=payload or {},
            headers={"X-ApiKey": api_key, "Content-Type": "application/json"},
            timeout=20,
        )
    except requests.exceptions.SSLError as e:
        return None, f"SSL certificate error connecting to {base_url}: {e}"
    except requests.exceptions.ConnectionError as e:
        return None, f"Couldn't connect to {base_url}: {e}"
    except requests.exceptions.Timeout:
        return None, f"Request to {base_url}{path} timed out."
    except requests.exceptions.RequestException as e:
        return None, f"Request failed: {e}"

    if resp.status_code == 401 or resp.status_code == 403:
        try:
            detail = resp.json()
        except ValueError:
            detail = resp.text.strip()
        # Airlock's own error body usually distinguishes an invalid key
        # from a key that's valid but missing a required REST API role
        # for this endpoint — surface it rather than a generic message.
        return None, f"Airlock Digital rejected the request (status {resp.status_code}): {detail}"

    log.info("Airlock %s -> %s: %s", path, resp.status_code, resp.text[:2000])

    try:
        data = resp.json()
    except ValueError:
        return None, f"Airlock Digital returned a non-JSON response (status {resp.status_code})."

    if not resp.ok:
        return None, f"Airlock Digital API error (status {resp.status_code}): {data}"

    error_field = data.get("error")
    if error_field is not None and str(error_field).strip().lower() != "success":
        return None, f"Airlock Digital API error: {error_field}"

    return data.get("response", {}), None


def airlock_move_agents(dest_groupid, agent_ids):
    """Moves the given agent IDs to dest_groupid. Returns (ok, error)."""
    if not agent_ids:
        return True, None
    _, err = airlock_request("/v1/agent/move", {"groupid": dest_groupid, "agentid": agent_ids})
    return err is None, err


def load_iso_mapping():
    """Loads the editable ISO 27001 <-> Airlock control mapping.
    Returns (mapping_dict, error) — mapping_dict is {} on any failure,
    and error explains specifically why (missing file vs. invalid JSON
    vs. an OS-level read error), rather than a single generic message,
    since those need different fixes."""
    filename = os.path.basename(ISO_MAPPING_PATH)
    if not os.path.exists(ISO_MAPPING_PATH):
        return {}, (
            f"{filename} not found next to app.py. Make sure it was copied into the same "
            "folder as app.py — it doesn't get replaced as often as the other files, so it's "
            "easy to leave behind when updating."
        )
    try:
        with open(ISO_MAPPING_PATH, "r") as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        return {}, f"{filename} contains invalid JSON: {e}"
    except OSError as e:
        return {}, f"Couldn't read {filename}: {e}"
    return data, None


# Statuses a rule can evaluate to. "unknown" means the field this rule
# depends on wasn't present in the policy response at all (e.g. an
# older Airlock version, or a typo in a hand-edited iso_mapping.json) —
# deliberately distinct from "unmet" so a report doesn't quietly claim
# a control failed when really the data just wasn't there to check.
STATUS_MEETS = "meets"
STATUS_PARTIAL = "partial"
STATUS_UNMET = "unmet"
STATUS_UNKNOWN = "unknown"
STATUS_ERROR = "error"

# Rank used to combine multiple sub-rule results in all_of/any_of —
# lower is "worse". Used to pick the overall status for a compound rule.
_STATUS_RANK = {STATUS_MEETS: 3, STATUS_PARTIAL: 2, STATUS_UNMET: 1, STATUS_UNKNOWN: 0, STATUS_ERROR: 0}


def evaluate_rule(rule, policy):
    """Evaluates one rule (a dict from iso_mapping.json) against a
    group's raw policy configuration dict (the 'response' object from
    /v1/group/policies). Returns (status, detail) where detail is a
    short human-readable explanation for the report/dashboard."""
    if not isinstance(rule, dict):
        return STATUS_ERROR, "Malformed rule."

    rule_type = rule.get("type")

    if rule_type == "manual":
        return STATUS_UNKNOWN, "Requires manual attestation — not automatically verifiable from policy data."

    if rule_type == "field_equals":
        field = rule.get("field")
        if field not in policy:
            return STATUS_UNKNOWN, f"Field '{field}' not present in policy response."
        value = policy.get(field)
        if value == rule.get("pass_value"):
            return STATUS_MEETS, f"{field} = {value}"
        if "partial_value" in rule and value == rule.get("partial_value"):
            return STATUS_PARTIAL, f"{field} = {value}"
        return STATUS_UNMET, f"{field} = {value}"

    if rule_type == "list_any_match":
        field = rule.get("field")
        if field not in policy:
            return STATUS_UNKNOWN, f"Field '{field}' not present in policy response."
        items = policy.get(field)
        if not isinstance(items, list):
            return STATUS_UNKNOWN, f"Field '{field}' is not a list."
        if not items:
            empty_status = rule.get("empty_status", STATUS_UNMET)
            return empty_status, f"'{field}' is empty."
        item_field = rule.get("item_field")
        match_value = rule.get("match_value")
        min_count = rule.get("min_count", 1)
        matches = sum(1 for it in items if isinstance(it, dict) and it.get(item_field) == match_value)
        if matches >= min_count:
            return STATUS_MEETS, f"{matches} of {len(items)} entr{'y' if len(items)==1 else 'ies'} in '{field}' match."
        return STATUS_PARTIAL, f"{matches} of {len(items)} entr{'y' if len(items)==1 else 'ies'} in '{field}' match (need {min_count})."

    if rule_type in ("all_of", "any_of"):
        sub_rules = rule.get("rules", [])
        if not sub_rules:
            return STATUS_ERROR, "Compound rule has no sub-rules."
        results = [evaluate_rule(r, policy) for r in sub_rules]
        statuses = [s for s, _ in results]
        details = "; ".join(d for _, d in results)
        if rule_type == "all_of":
            overall = min(statuses, key=lambda s: _STATUS_RANK[s])
        else:
            overall = max(statuses, key=lambda s: _STATUS_RANK[s])
        return overall, details

    return STATUS_ERROR, f"Unknown rule type '{rule_type}'."


def evaluate_group_against_mapping(policy, mapping):
    """Evaluates every control in the mapping against one group's raw
    policy configuration. Returns a list of {id, title, name, status,
    detail} dicts."""
    results = []
    for control in mapping.get("controls", []):
        status, detail = evaluate_rule(control.get("rule", {}), policy)
        results.append(
            {
                "id": control.get("id"),
                "title": control.get("title"),
                "name": control.get("name"),
                "theme": control.get("theme"),
                "status": status,
                "detail": detail,
            }
        )
    return results


# --- Custom Widgets: Partner Engagement Report (Cloud multi-tenant API) ---
#
# This talks to a completely different API surface than the rest of the
# app — Airlock Digital's Cloud/MSP management layer, not a single
# on-prem server. Auth uses a "UserApiKey" header (not "X-ApiKey" like
# the on-prem API), plus per-tenant "tenantID" and "Directoryid"
# headers. URLs follow https://<base_domain>/<module>/v1/<endpoint>,
# where the module segment varies by feature area (confirmed via live
# testing: "webfe" for most tenant data, "policy" for the policy list).
#
# Every endpoint/field name below was confirmed against real API
# responses during development — see project history for the captured
# samples this was built from.

CLOUD_MODULE_WEBFE = "webfe"
CLOUD_MODULE_POLICY = "policy"
CLOUD_MODULE_DIRECTORY = "directory"


def get_cloud_admin_config():
    cfg = load_cloud_config()
    return cfg.get("cloud_admin") or {}


def cloud_api_request(module, endpoint, method="GET", payload=None, tenant_id="", directory_id=""):
    """Calls the Cloud multi-tenant API. Returns (data, error)."""
    cfg = get_cloud_admin_config()
    base_domain = (cfg.get("base_domain") or "").strip()
    api_key = (cfg.get("api_key") or "").strip()

    if not base_domain:
        return None, "No Cloud base domain configured. Add one in the Partner Engagement Report widget."
    if not api_key:
        return None, "No Cloud admin API key configured. Add one in the Partner Engagement Report widget."

    url = f"https://{base_domain}/{module}/v1/{endpoint}"
    headers = {
        "Content-Type": "application/json",
        "UserApiKey": api_key,
        "tenantID": tenant_id,
        "Directoryid": directory_id,
    }

    try:
        if method == "GET":
            resp = requests.get(url, headers=headers, timeout=20)
        else:
            resp = requests.post(url, json=payload or {}, headers=headers, timeout=20)
    except requests.exceptions.SSLError as e:
        return None, f"SSL certificate error connecting to {url}: {e}"
    except requests.exceptions.ConnectionError as e:
        return None, f"Couldn't connect to {url}: {e}"
    except requests.exceptions.Timeout:
        return None, f"Request to {url} timed out."
    except requests.exceptions.RequestException as e:
        return None, f"Request failed: {e}"

    log.info("Cloud API %s %s (tenant=%s) -> %s: %s", method, url, tenant_id, resp.status_code, resp.text[:1000])

    if resp.status_code in (401, 403):
        try:
            detail = resp.json()
        except ValueError:
            detail = resp.text.strip()
        return None, f"Cloud API rejected the request (status {resp.status_code}): {detail}"

    if resp.status_code == 405:
        return None, f"Cloud API rejected the HTTP method for {module}/v1/{endpoint} (405 Method Not Allowed) — tried {method}."

    try:
        data = resp.json()
    except ValueError:
        return None, f"Cloud API returned a non-JSON response (status {resp.status_code}): {resp.text[:300]}"

    if not resp.ok:
        return None, f"Cloud API error (status {resp.status_code}): {data}"

    return data, None


def cloud_get_tenant_users(tenant_id, directory_id):
    data, err = cloud_api_request(CLOUD_MODULE_WEBFE, "tenant-user-list", tenant_id=tenant_id, directory_id=directory_id)
    if err:
        return None, err
    users = (((data or {}).get("TenantUserList") or {}).get("TenantUserList")) or []
    result = []
    for u in users:
        if not isinstance(u, dict):
            continue
        result.append(
            {
                "full_name": u.get("FullName") or f"{u.get('FirstName', '')} {u.get('LastName', '')}".strip(),
                "email": u.get("Email", ""),
                "last_accessed": u.get("LastAccessedDate", ""),
            }
        )
    return result, None


def cloud_get_managed_count(tenant_id, directory_id):
    # Confirmed via live testing: this one needs POST (with an empty
    # body), unlike tenant-user-list which needs GET — no consistent
    # rule across endpoints, so this has to be tracked per-endpoint.
    data, err = cloud_api_request(
        CLOUD_MODULE_WEBFE,
        "tenant-policy-managed-count-list",
        method="POST",
        payload={},
        tenant_id=tenant_id,
        directory_id=directory_id,
    )
    if err:
        return None, err
    tcc = (((data or {}).get("TenantPolicyManagedCountList") or {}).get("TenantClientCount")) or {}
    return {"client_count": tcc.get("ClientCount", 0), "unmanaged_count": tcc.get("UnmanagedCount", 0)}, None


def cloud_get_policies(tenant_id, directory_id):
    data, err = cloud_api_request(
        CLOUD_MODULE_POLICY, "policy?limit=10000", tenant_id=tenant_id, directory_id=directory_id
    )
    if err:
        return None, err
    policies = (data or {}).get("data") or []
    result = []
    for p in policies:
        if not isinstance(p, dict):
            continue
        result.append(
            {
                "id": p.get("policyid") or p.get("id"),
                "name": p.get("name", ""),
                "auditmode": str(p.get("auditmode", "")),  # "0" = enforce, "1" = audit
            }
        )
    return result, None


def cloud_get_policy_client_count(tenant_id, directory_id, policy_id, policy_name):
    payload = {"TenantPolicyGroupList": [{"ID": policy_id, "Name": policy_name}], "atcp": None}
    data, err = cloud_api_request(
        CLOUD_MODULE_WEBFE,
        "tenant-policy-clients-in-scope-list",
        method="POST",
        payload=payload,
        tenant_id=tenant_id,
        directory_id=directory_id,
    )
    if err:
        return None, err
    selected = (data or {}).get("TenantClientSelectedList") or []
    if selected and isinstance(selected[0], dict):
        return selected[0].get("InScope", 0), None
    return 0, None


def cloud_get_license_allocation(tenant_id, directory_id):
    payload = {"DirectoryLicenceAllocationList": {"DirectoryID": directory_id}, "atcp": None}
    data, err = cloud_api_request(
        CLOUD_MODULE_DIRECTORY,
        "directory-licence-allocation-list",
        method="POST",
        payload=payload,
        tenant_id=tenant_id,
        directory_id=directory_id,
    )
    if err:
        return None, err
    dlal = (data or {}).get("DirectoryLicenceAllocationList") or {}
    return dlal.get("DirectoryLicenceAllocated", 0), None


def cloud_collect_tenant_data(tenant_entry):
    """Pulls all four data points for one saved partner tenant. Never
    raises, and each data point is independent — a failure in one
    (e.g. license data) is recorded on its own error field rather than
    blanking out the whole tenant, since the four calls have nothing to
    do with each other."""
    tenant_id = tenant_entry.get("tenant_id", "")
    directory_id = tenant_entry.get("directory_id", "")
    label = tenant_entry.get("label", "")

    result = {
        "label": label,
        "tenant_id": tenant_id,
        "directory_id": directory_id,
        "users": [],
        "users_error": None,
        "client_count": 0,
        "unmanaged_count": 0,
        "agents_error": None,
        "audit_count": 0,
        "enforce_count": 0,
        "policy_errors": [],
        "license_allocated": None,
        "license_error": None,
    }

    users, err = cloud_get_tenant_users(tenant_id, directory_id)
    if err:
        result["users_error"] = err
    else:
        result["users"] = users

    counts, err = cloud_get_managed_count(tenant_id, directory_id)
    if err:
        result["agents_error"] = err
    else:
        result["client_count"] = counts["client_count"]
        result["unmanaged_count"] = counts["unmanaged_count"]

    policies, err = cloud_get_policies(tenant_id, directory_id)
    if err:
        result["policy_errors"].append(f"Couldn't load the policy list: {err}")
    else:
        audit_total = 0
        enforce_total = 0
        for p in policies:
            in_scope, perr = cloud_get_policy_client_count(tenant_id, directory_id, p["id"], p["name"])
            if perr:
                result["policy_errors"].append(f"{p['name']}: {perr}")
                continue
            if p["auditmode"] == "0":
                enforce_total += in_scope or 0
            else:
                audit_total += in_scope or 0
        result["audit_count"] = audit_total
        result["enforce_count"] = enforce_total

    license_allocated, err = cloud_get_license_allocation(tenant_id, directory_id)
    if err:
        result["license_error"] = err
    else:
        result["license_allocated"] = license_allocated

    return result


def run_cloud_partner_report(tenant_ids=None):
    cloud_cfg = load_cloud_config()
    tenants = cloud_cfg.get("cloud_tenants") or []
    if tenant_ids is not None:
        tenants = [t for t in tenants if t.get("id") in tenant_ids]
    if not tenants:
        return None, "No partner tenants selected. Choose at least one and try again."

    results = [cloud_collect_tenant_data(t) for t in tenants]
    return {"tenants": results, "generated_at": datetime.now().isoformat()}, None


def parse_github_repo(value):
    """Accepts 'owner/repo' or a full https://github.com/owner/repo URL."""
    value = value.strip()
    if value.endswith(".git"):
        value = value[: -len(".git")]
    for prefix in ("https://github.com/", "http://github.com/", "github.com/"):
        if value.startswith(prefix):
            value = value[len(prefix):]
            break
    value = value.strip("/")
    parts = value.split("/")
    if len(parts) < 2 or not parts[0] or not parts[1]:
        raise ValueError("Repo must be in 'owner/repo' form, or a full GitHub URL.")
    return parts[0], parts[1]


def sync_github_scripts():
    """Downloads .py files from the configured GitHub repo into the local
    scripts folder. Only touches .py files it manages — any other files
    already sitting in the scripts folder (notes, configs, etc.) are left
    alone. Returns a dict describing the outcome; never raises.
    """
    config = load_config()
    if (config.get("github_sync_enabled", "") or "").strip().lower() != "true":
        return {"synced": False, "reason": "disabled"}

    repo = config.get("github_repo", "").strip()
    if not repo:
        return {"synced": False, "reason": "no repo configured"}

    branch = config.get("github_branch", "").strip() or "main"
    subdir = config.get("github_subdir", "").strip().strip("/")
    token = config.get("github_token", "").strip()

    try:
        owner, name = parse_github_repo(repo)
    except ValueError as e:
        return {"synced": False, "reason": str(e)}

    api_url = f"https://api.github.com/repos/{owner}/{name}/contents/{subdir}"
    headers = {"Accept": "application/vnd.github.v3+json"}
    if token:
        headers["Authorization"] = f"token {token}"

    try:
        resp = requests.get(api_url, headers=headers, params={"ref": branch}, timeout=15)
    except requests.exceptions.RequestException as e:
        return {"synced": False, "reason": f"GitHub request failed: {e}"}

    if not resp.ok:
        return {"synced": False, "reason": f"GitHub API error: {resp.status_code}"}

    items = resp.json()
    if not isinstance(items, list):
        return {"synced": False, "reason": "Configured path isn't a folder in this repo."}

    py_files = [item for item in items if item.get("type") == "file" and item["name"].endswith(".py")]

    scripts_dir = get_scripts_dir()
    downloaded = []
    errors = []
    for item in py_files:
        raw_url = item.get("download_url")
        if not raw_url:
            continue
        try:
            file_resp = requests.get(raw_url, headers=headers, timeout=15)
            file_resp.raise_for_status()
        except requests.exceptions.RequestException as e:
            errors.append(f"{item['name']}: {e}")
            continue
        with open(os.path.join(scripts_dir, item["name"]), "wb") as f:
            f.write(file_resp.content)
        downloaded.append(item["name"])

    return {
        "synced": True,
        "downloaded": downloaded,
        "errors": errors,
        "location": f"{owner}/{name}@{branch}" + (f"/{subdir}" if subdir else ""),
    }


# --- Settings: Publish to GitHub ---
# The other direction from GitHub sync above — pushes this app's own
# source *out* to a repo instead of pulling scripts *in*. Deliberately a
# fixed, hand-maintained file list rather than a real .gitignore parser:
# this app has one, small, known set of source files, and being explicit
# here means a secret or runtime file (config.json, cloud_config.json,
# venv/, builds/, audit_sessions.json, logs) can never accidentally end
# up on the list just because someone dropped a new file next to app.py.
GITHUB_PUBLISH_FILES = [
    ".gitignore",
    "app.py",
    "iso_mapping.json",
    "PROJECT_SUMMARY.md",
    "README.md",
    "requirements.txt",
    "run.bat",
    os.path.join("templates", "index.html"),
]


def collect_publish_files():
    """Returns (found, missing) where found is [(repo_path, abs_path), ...]
    for files in GITHUB_PUBLISH_FILES that currently exist on disk, and
    missing is the repo-relative paths that don't (e.g. no PROJECT_SUMMARY.md
    yet) — not fatal, just left out and reported back."""
    found = []
    missing = []
    for rel in GITHUB_PUBLISH_FILES:
        abs_path = os.path.join(BASE_DIR, rel)
        repo_path = rel.replace(os.sep, "/")
        if os.path.isfile(abs_path):
            found.append((repo_path, abs_path))
        else:
            missing.append(repo_path)
    return found, missing


def github_publish_push(commit_message=None):
    """Pushes GITHUB_PUBLISH_FILES to the configured GitHub repo as a
    single commit, using the Git Data API directly (blobs -> tree ->
    commit -> ref update) rather than the simpler per-file Contents API,
    so a multi-file update lands as one commit instead of one per file.
    Never raises — returns (result_dict, None) or (None, error_message)."""
    config = load_config()
    repo = config.get("github_publish_repo", "").strip()
    if not repo:
        return None, "No publish repo configured."
    branch = config.get("github_publish_branch", "").strip() or "main"
    token = config.get("github_publish_token", "").strip()
    if not token:
        return None, "A personal access token with write access to the repo is required."

    try:
        owner, name = parse_github_repo(repo)
    except ValueError as e:
        return None, str(e)

    files, missing = collect_publish_files()
    if not files:
        return None, "None of the expected project files were found on disk."

    headers = {
        "Accept": "application/vnd.github+json",
        "Authorization": f"Bearer {token}",
    }
    api_base = f"https://api.github.com/repos/{owner}/{name}"

    # Look up the branch's current tip commit, if the branch already
    # exists — a brand-new repo/branch won't have one yet, and that's
    # fine: we just create the first commit with no parent.
    parent_sha = None
    base_tree_sha = None
    try:
        ref_resp = requests.get(f"{api_base}/git/ref/heads/{branch}", headers=headers, timeout=15)
    except requests.exceptions.RequestException as e:
        return None, f"GitHub request failed: {e}"

    if ref_resp.status_code == 200:
        parent_sha = ref_resp.json()["object"]["sha"]
        try:
            commit_resp = requests.get(f"{api_base}/git/commits/{parent_sha}", headers=headers, timeout=15)
            commit_resp.raise_for_status()
            base_tree_sha = commit_resp.json()["tree"]["sha"]
        except requests.exceptions.RequestException as e:
            return None, f"Couldn't read the branch's current commit: {e}"
    elif ref_resp.status_code != 404:
        return None, f"GitHub API error looking up branch: {ref_resp.status_code} {ref_resp.text[:200]}"

    # Upload each file as a blob first — the tree just references their SHAs.
    tree_entries = []
    for repo_path, abs_path in files:
        try:
            with open(abs_path, "rb") as f:
                content_bytes = f.read()
        except OSError as e:
            return None, f"Couldn't read {repo_path}: {e}"

        blob_payload = {"content": base64.b64encode(content_bytes).decode("ascii"), "encoding": "base64"}
        try:
            blob_resp = requests.post(f"{api_base}/git/blobs", headers=headers, json=blob_payload, timeout=30)
        except requests.exceptions.RequestException as e:
            return None, f"GitHub request failed uploading {repo_path}: {e}"
        if not blob_resp.ok:
            return None, f"GitHub API error uploading {repo_path}: {blob_resp.status_code} {blob_resp.text[:200]}"

        tree_entries.append({
            "path": repo_path,
            "mode": "100644",
            "type": "blob",
            "sha": blob_resp.json()["sha"],
        })

    tree_payload = {"tree": tree_entries}
    if base_tree_sha:
        tree_payload["base_tree"] = base_tree_sha
    try:
        tree_resp = requests.post(f"{api_base}/git/trees", headers=headers, json=tree_payload, timeout=30)
    except requests.exceptions.RequestException as e:
        return None, f"GitHub request failed creating tree: {e}"
    if not tree_resp.ok:
        return None, f"GitHub API error creating tree: {tree_resp.status_code} {tree_resp.text[:200]}"
    new_tree_sha = tree_resp.json()["sha"]

    message = (commit_message or "").strip() or f"Update from Partner Consulting Toolkit ({datetime.now().strftime('%Y-%m-%d %H:%M')})"
    commit_payload = {"message": message, "tree": new_tree_sha}
    if parent_sha:
        commit_payload["parents"] = [parent_sha]
    try:
        new_commit_resp = requests.post(f"{api_base}/git/commits", headers=headers, json=commit_payload, timeout=30)
    except requests.exceptions.RequestException as e:
        return None, f"GitHub request failed creating commit: {e}"
    if not new_commit_resp.ok:
        return None, f"GitHub API error creating commit: {new_commit_resp.status_code} {new_commit_resp.text[:200]}"
    new_commit_sha = new_commit_resp.json()["sha"]

    # Existing branch: fast-forward its ref. Brand-new branch: create the ref.
    if parent_sha:
        try:
            update_resp = requests.patch(
                f"{api_base}/git/refs/heads/{branch}", headers=headers, json={"sha": new_commit_sha}, timeout=15
            )
        except requests.exceptions.RequestException as e:
            return None, f"GitHub request failed updating branch ref: {e}"
    else:
        try:
            update_resp = requests.post(
                f"{api_base}/git/refs",
                headers=headers,
                json={"ref": f"refs/heads/{branch}", "sha": new_commit_sha},
                timeout=15,
            )
        except requests.exceptions.RequestException as e:
            return None, f"GitHub request failed creating branch ref: {e}"
    if not update_resp.ok:
        return None, f"GitHub API error updating branch: {update_resp.status_code} {update_resp.text[:200]}"

    return {
        "commit_sha": new_commit_sha,
        "commit_url": f"https://github.com/{owner}/{name}/commit/{new_commit_sha}",
        "files_pushed": [p for p, _ in files],
        "files_missing": missing,
        "branch": branch,
    }, None


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/run", methods=["POST"])
def run_code():
    data = request.get_json(silent=True) or {}
    code = data.get("code", "")

    if not isinstance(code, str) or not code.strip():
        return jsonify({"error": "No code provided."}), 400

    # Write the submitted code to a temporary file and run it as a
    # separate process, so it can't crash or hang the server itself.
    tmp_dir = tempfile.gettempdir()
    filename = os.path.join(tmp_dir, f"snippet_{uuid.uuid4().hex}.py")

    try:
        with open(filename, "w") as f:
            f.write(code)

        try:
            result = subprocess.run(
                [sys.executable, filename],
                capture_output=True,
                text=True,
                timeout=EXECUTION_TIMEOUT,
            )
            return jsonify(
                {
                    "stdout": result.stdout,
                    "stderr": result.stderr,
                    "exit_code": result.returncode,
                    "timed_out": False,
                }
            )
        except subprocess.TimeoutExpired as e:
            return jsonify(
                {
                    "stdout": e.stdout or "",
                    "stderr": (e.stderr or "") + f"\n[Execution stopped: exceeded {EXECUTION_TIMEOUT}s time limit]",
                    "exit_code": None,
                    "timed_out": True,
                }
            )
    finally:
        if os.path.exists(filename):
            os.remove(filename)


@app.route("/compile", methods=["POST"])
def compile_code():
    data = request.get_json(silent=True) or {}
    code = data.get("code", "")

    if not isinstance(code, str) or not code.strip():
        return jsonify({"success": False, "error": "No code provided."}), 400

    # Name the output using today's date and time, e.g. script_20260714_143205.exe
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    exe_basename = f"script_{timestamp}"

    with tempfile.TemporaryDirectory() as work_dir:
        script_path = os.path.join(work_dir, "source.py")
        with open(script_path, "w") as f:
            f.write(code)

        cmd = [
            sys.executable, "-m", "PyInstaller",
            "--onefile",
            "--noconfirm",
            "--distpath", BUILDS_DIR,
            "--workpath", work_dir,
            "--specpath", work_dir,
            "--name", exe_basename,
            script_path,
        ]

        try:
            result = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=COMPILE_TIMEOUT,
            )
        except FileNotFoundError:
            return jsonify(
                {
                    "success": False,
                    "error": "PyInstaller isn't installed. Run: pip install pyinstaller",
                }
            ), 500
        except subprocess.TimeoutExpired:
            return jsonify(
                {
                    "success": False,
                    "error": f"Compilation timed out after {COMPILE_TIMEOUT}s.",
                }
            ), 500

        # PyInstaller only appends .exe automatically on Windows.
        produced_name = exe_basename + (".exe" if os.name == "nt" else "")
        produced_path = os.path.join(BUILDS_DIR, produced_name)
        built_ok = result.returncode == 0 and os.path.exists(produced_path)

        return jsonify(
            {
                "success": built_ok,
                "filename": produced_name if built_ok else None,
                "download_url": f"/download/{produced_name}" if built_ok else None,
                # Keep logs short — PyInstaller output can be very long.
                "stdout": result.stdout[-3000:],
                "stderr": result.stderr[-3000:],
                "exit_code": result.returncode,
            }
        )


@app.route("/download/<path:filename>")
def download(filename):
    return send_from_directory(BUILDS_DIR, filename, as_attachment=True)


def key_preview(value):
    return ("•" * 8 + value[-4:]) if value else ""


@app.route("/config", methods=["GET"])
def get_config():
    config = load_config()
    vt_key = config.get("vt_api_key", "")
    github_token = config.get("github_token", "")
    github_publish_token = config.get("github_publish_token", "")

    active_profile = get_active_airlock_profile(config)
    airlock_key = (active_profile.get("api_key") if active_profile else "") or ""

    profiles = config.get("airlock_profiles") or []
    profiles_out = [
        {
            "id": p.get("id"),
            "label": p.get("label", ""),
            "tenant": p.get("tenant", ""),
            "port": p.get("port", "") or str(DEFAULT_AIRLOCK_PORT),
            "has_key": bool(p.get("api_key")),
            "key_preview": key_preview(p.get("api_key", "")),
        }
        for p in profiles
    ]

    return jsonify(
        {
            "vt_api_key": {"has_key": bool(vt_key), "key_preview": key_preview(vt_key)},
            # Reflects the active Airlock connection (see airlock_profiles
            # below for the full saved list) — kept under these names for
            # backward compatibility with the Scripts tab's info display.
            "airlock_api_key": {"has_key": bool(airlock_key), "key_preview": key_preview(airlock_key)},
            "airlock_tenant": (active_profile.get("tenant") if active_profile else "") or "",
            "airlock_port": (active_profile.get("port") if active_profile else "") or "",
            "airlock_profiles": profiles_out,
            "active_airlock_profile_id": config.get("active_airlock_profile_id"),
            "active_airlock_profile_label": (active_profile.get("label") if active_profile else "") or "",
            "api_scripts_dir": config.get("api_scripts_dir", ""),
            "api_scripts_dir_resolved": get_scripts_dir(),
            "api_script_timeout": config.get("api_script_timeout", ""),
            "api_script_timeout_resolved": get_script_timeout(),
            "github_sync_enabled": (config.get("github_sync_enabled", "") or "").strip().lower() == "true",
            "github_repo": config.get("github_repo", ""),
            "github_branch": config.get("github_branch", ""),
            "github_subdir": config.get("github_subdir", ""),
            "github_token": {"has_key": bool(github_token), "key_preview": key_preview(github_token)},
            "github_publish_repo": config.get("github_publish_repo", ""),
            "github_publish_branch": config.get("github_publish_branch", ""),
            "github_publish_token": {"has_key": bool(github_publish_token), "key_preview": key_preview(github_publish_token)},
            # Defaults to true (on) when unset, matching this app's
            # behavior before the toggle existed.
            "logging_enabled": (config.get("logging_enabled") if config.get("logging_enabled") is not None else "true").strip().lower() == "true",
        }
    )


VALID_CONFIG_FIELDS = {
    "vt_api_key", "api_scripts_dir", "api_script_timeout",
    "github_sync_enabled", "github_repo", "github_branch", "github_subdir", "github_token",
    "github_publish_repo", "github_publish_branch", "github_publish_token",
    "logging_enabled",
}
SECRET_CONFIG_FIELDS = {"vt_api_key", "github_token", "github_publish_token"}


@app.route("/config", methods=["POST"])
def set_config():
    data = request.get_json(silent=True) or {}

    # Supports either a single {"field": ..., "value": ...} update, or a
    # batch {"fields": {"a": "...", "b": "..."}} update in one request.
    if "fields" in data and isinstance(data["fields"], dict):
        updates = data["fields"]
    elif "field" in data:
        updates = {data.get("field"): data.get("value", "")}
    else:
        return jsonify({"error": "No fields provided."}), 400

    if not updates:
        return jsonify({"error": "No fields provided."}), 400

    for field, value in updates.items():
        if field not in VALID_CONFIG_FIELDS:
            return jsonify({"error": f"Unknown field '{field}'."}), 400
        if not isinstance(value, str):
            return jsonify({"error": f"Value for '{field}' must be a string."}), 400

    if "github_sync_enabled" in updates and updates["github_sync_enabled"].strip().lower() not in ("true", "false"):
        return jsonify({"error": "github_sync_enabled must be 'true' or 'false'."}), 400

    if "logging_enabled" in updates and updates["logging_enabled"].strip().lower() not in ("true", "false"):
        return jsonify({"error": "logging_enabled must be 'true' or 'false'."}), 400

    if "api_script_timeout" in updates:
        timeout_val = updates["api_script_timeout"].strip()
        if timeout_val and not (timeout_val.isdigit() and int(timeout_val) > 0):
            return jsonify({"error": "api_script_timeout must be a whole number of seconds greater than 0."}), 400

    config = load_config()
    for field, value in updates.items():
        config[field] = value.strip()
    save_config(config)

    # Apply immediately — no restart needed to see the effect.
    if "logging_enabled" in updates:
        configure_logging(updates["logging_enabled"].strip().lower() == "true")

    response = {"saved": True}
    if len(updates) == 1:
        only_field = next(iter(updates))
        response["field"] = only_field
        if only_field in SECRET_CONFIG_FIELDS:
            response["key_preview"] = key_preview(config[only_field])
        if only_field == "api_scripts_dir":
            response["api_scripts_dir_resolved"] = get_scripts_dir()
        if only_field == "api_script_timeout":
            response["api_script_timeout_resolved"] = get_script_timeout()
    else:
        response["fields"] = list(updates.keys())
        secret_updates = {f: key_preview(config[f]) for f in updates if f in SECRET_CONFIG_FIELDS}
        if secret_updates:
            response["key_previews"] = secret_updates
        if "api_scripts_dir" in updates:
            response["api_scripts_dir_resolved"] = get_scripts_dir()
        if "api_script_timeout" in updates:
            response["api_script_timeout_resolved"] = get_script_timeout()

    return jsonify(response)


@app.route("/config/reveal", methods=["POST"])
def reveal_config_value():
    # Returns a saved secret's real value so the Settings tab can copy it
    # to the clipboard. Only ever called by an explicit user click on a
    # Copy button — never included in the regular masked /config GET.
    data = request.get_json(silent=True) or {}
    field = data.get("field", "")

    if field not in SECRET_CONFIG_FIELDS:
        return jsonify({"error": "Invalid field."}), 400

    config = load_config()
    value = config.get(field, "")
    if not value:
        return jsonify({"error": "No key saved."}), 404

    return jsonify({"value": value})


@app.route("/api_scripts", methods=["GET"])
def list_api_scripts():
    scripts_dir = get_scripts_dir()
    try:
        scripts = sorted(
            f for f in os.listdir(scripts_dir)
            if f.endswith(".py") and os.path.isfile(os.path.join(scripts_dir, f))
        )
    except OSError as e:
        return jsonify({"error": f"Couldn't read scripts folder: {e}", "scripts": [], "location": scripts_dir}), 500
    return jsonify({"scripts": scripts, "location": scripts_dir})


@app.route("/sync", methods=["POST"])
def sync_now():
    result = sync_github_scripts()
    if not result["synced"]:
        return jsonify({"success": False, "error": result.get("reason", "Sync failed.")}), 400
    return jsonify(
        {
            "success": True,
            "downloaded": result["downloaded"],
            "errors": result["errors"],
            "location": result["location"],
        }
    )


@app.route("/github/publish/files", methods=["GET"])
def github_publish_files():
    found, missing = collect_publish_files()
    return jsonify({"files": [p for p, _ in found], "missing": missing})


@app.route("/github/publish", methods=["POST"])
def github_publish():
    data_in = request.get_json(silent=True) or {}
    commit_message = data_in.get("message", "")
    if not isinstance(commit_message, str):
        return jsonify({"error": "message must be a string."}), 400

    result, err = github_publish_push(commit_message=commit_message)
    if err:
        return jsonify({"error": err}), 400
    return jsonify({"success": True, **result})


@app.route("/api_scripts/run", methods=["POST"])
def run_api_script():
    data = request.get_json(silent=True) or {}
    # "script" holds the script name optionally followed by command-line
    # arguments, e.g. "report.py --user jdoe --verbose".
    raw_input_str = data.get("script", "")

    if not isinstance(raw_input_str, str) or not raw_input_str.strip():
        return jsonify({"error": "No script selected."}), 400

    try:
        tokens = shlex.split(raw_input_str.strip())
    except ValueError as e:
        return jsonify({"error": f"Couldn't parse arguments: {e}"}), 400

    if not tokens:
        return jsonify({"error": "No script selected."}), 400

    script_name, script_args = tokens[0], tokens[1:]

    safe_name = os.path.basename(script_name)
    if not safe_name.endswith(".py"):
        return jsonify({"error": "Invalid script name."}), 400

    exec_dir = get_scripts_dir()
    script_path = os.path.join(exec_dir, safe_name)
    if not os.path.isfile(script_path):
        return jsonify({"error": f"Script not found: {safe_name}"}), 404

    config = load_config()
    active_profile = get_active_airlock_profile(config)
    env = os.environ.copy()
    env["AIRLOCK_API_KEY"] = (active_profile.get("api_key") if active_profile else "") or ""
    env["AIRLOCK_TENANT"] = (active_profile.get("tenant") if active_profile else "") or ""
    env["AIRLOCK_PORT"] = ((active_profile.get("port") if active_profile else "") or "") or str(DEFAULT_AIRLOCK_PORT)
    env["VT_API_KEY"] = config.get("vt_api_key", "")

    # Snapshot existing xlsx/xml/html files beforehand (including
    # subfolders, since scripts often write into a folder they create)
    # so we can tell which one the script produced or updated.
    before_mtimes = {}
    try:
        for root, _dirs, files in os.walk(exec_dir):
            for f in files:
                if f.lower().endswith(OUTPUT_FILE_EXTENSIONS):
                    fp = os.path.join(root, f)
                    before_mtimes[fp] = os.path.getmtime(fp)
    except OSError:
        pass

    run_start = time.time()
    script_timeout = get_script_timeout()

    try:
        result = subprocess.run(
            [sys.executable, script_path, *script_args],
            capture_output=True,
            text=True,
            timeout=script_timeout,
            cwd=exec_dir,
            env=env,
        )
        response = {
            "stdout": result.stdout,
            "stderr": result.stderr,
            "exit_code": result.returncode,
            "timed_out": False,
        }
    except subprocess.TimeoutExpired as e:
        response = {
            "stdout": e.stdout or "",
            "stderr": (e.stderr or "") + f"\n[Execution stopped: exceeded {script_timeout}s time limit]",
            "exit_code": None,
            "timed_out": True,
        }

    # Detect the output file, preferring a path the script explicitly
    # announced in its own console output (most reliable — works even
    # when the file lives in a subfolder the script just created).
    output_file = extract_output_path_from_text(response.get("stdout", ""), exec_dir)
    if not output_file:
        output_file = extract_output_path_from_text(response.get("stderr", ""), exec_dir)

    if not output_file:
        # Fall back to spotting a new or freshly-modified xlsx/xml/html
        # file anywhere under the scripts folder, including subfolders.
        try:
            candidates = []
            for root, _dirs, files in os.walk(exec_dir):
                for f in files:
                    if not f.lower().endswith(OUTPUT_FILE_EXTENSIONS):
                        continue
                    fp = os.path.join(root, f)
                    mtime = os.path.getmtime(fp)
                    is_new_or_updated = fp not in before_mtimes or mtime > before_mtimes[fp]
                    if is_new_or_updated and mtime >= run_start - 1:
                        candidates.append((mtime, fp))
            if candidates:
                candidates.sort(reverse=True)
                output_file = candidates[0][1]
        except OSError:
            pass

    if output_file:
        response["output_file"] = output_file
        response["output_filename"] = os.path.basename(output_file)

    return jsonify(response)


@app.route("/open_file", methods=["POST"])
def open_file():
    data = request.get_json(silent=True) or {}
    path = data.get("path", "")

    if not isinstance(path, str) or not path.strip():
        return jsonify({"success": False, "error": "No path provided."}), 400

    real_path = os.path.realpath(path)
    allowed_roots = [os.path.realpath(get_scripts_dir())]
    if not any(real_path == root or real_path.startswith(root + os.sep) for root in allowed_roots):
        return jsonify({"success": False, "error": "That file is outside the allowed scripts/output folders."}), 403

    if not os.path.isfile(real_path):
        return jsonify({"success": False, "error": "File not found."}), 404

    try:
        if os.name == "nt":
            os.startfile(real_path)  # noqa: os.startfile only exists on Windows
        elif sys.platform == "darwin":
            subprocess.run(["open", real_path], check=False)
        else:
            subprocess.run(["xdg-open", real_path], check=False)
        return jsonify({"success": True})
    except OSError as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/virustotal/upload", methods=["POST"])
def virustotal_upload():
    data = request.get_json(silent=True) or {}
    filename = data.get("filename", "")

    if not isinstance(filename, str) or not filename.strip():
        return jsonify({"success": False, "error": "No filename provided."}), 400

    # Prevent path traversal — only allow files that actually live in BUILDS_DIR.
    safe_name = os.path.basename(filename)
    file_path = os.path.join(BUILDS_DIR, safe_name)
    if not os.path.exists(file_path):
        return jsonify({"success": False, "error": f"File not found: {safe_name}"}), 404

    api_key = load_config().get("vt_api_key", "")
    if not api_key:
        return jsonify(
            {"success": False, "error": "No VirusTotal API key configured. Add one in the Settings tab."}
        ), 400

    headers = {"x-apikey": api_key}
    file_size = os.path.getsize(file_path)

    try:
        if file_size > VT_LARGE_FILE_THRESHOLD:
            # Files over 32MB need a dedicated upload URL first.
            resp = requests.get(VT_UPLOAD_URL_ENDPOINT, headers=headers, timeout=30)
            resp.raise_for_status()
            upload_url = resp.json()["data"]
        else:
            upload_url = VT_FILES_URL

        with open(file_path, "rb") as f:
            resp = requests.post(
                upload_url,
                headers=headers,
                files={"file": (safe_name, f, "application/octet-stream")},
                timeout=120,
            )

        if resp.status_code == 401:
            return jsonify({"success": False, "error": "VirusTotal rejected the API key (unauthorized)."}), 401
        if resp.status_code == 429:
            return jsonify({"success": False, "error": "VirusTotal rate limit / quota exceeded."}), 429
        resp.raise_for_status()

        analysis_id = resp.json()["data"]["id"]
        vt_url = f"https://www.virustotal.com/gui/file-analysis/{analysis_id}"

        return jsonify({"success": True, "analysis_id": analysis_id, "vt_url": vt_url})

    except requests.exceptions.RequestException as e:
        return jsonify({"success": False, "error": f"VirusTotal request failed: {e}"}), 502


# --- Airlock Digital connections (multiple tenant/key pairs) ---

@app.route("/airlock/profiles", methods=["GET"])
def airlock_profiles_list():
    config = load_config()
    profiles = config.get("airlock_profiles") or []
    result = [
        {
            "id": p.get("id"),
            "label": p.get("label", ""),
            "tenant": p.get("tenant", ""),
            "port": p.get("port", "") or str(DEFAULT_AIRLOCK_PORT),
            "has_key": bool(p.get("api_key")),
            "key_preview": key_preview(p.get("api_key", "")),
        }
        for p in profiles
    ]
    return jsonify({"profiles": result, "active_id": config.get("active_airlock_profile_id")})


@app.route("/airlock/profiles", methods=["POST"])
def airlock_profiles_add():
    data_in = request.get_json(silent=True) or {}
    label = (data_in.get("label") or "").strip()
    tenant = (data_in.get("tenant") or "").strip()
    port = (data_in.get("port") or "").strip()
    api_key = (data_in.get("api_key") or "").strip()

    if not label:
        return jsonify({"error": "A label is required."}), 400
    if not tenant:
        return jsonify({"error": "A tenant is required."}), 400
    if not api_key:
        return jsonify({"error": "An API key is required."}), 400
    if port and not (port.isdigit() and 1 <= int(port) <= 65535):
        return jsonify({"error": "Port must be a whole number between 1 and 65535."}), 400

    config = load_config()
    profiles = config.get("airlock_profiles") or []

    if any(p.get("label", "").strip().lower() == label.lower() for p in profiles):
        return jsonify({"error": f"A connection named '{label}' already exists."}), 400

    profile = {"id": str(uuid.uuid4()), "label": label, "tenant": tenant, "port": port, "api_key": api_key}
    profiles.append(profile)
    config["airlock_profiles"] = profiles
    # If this is the very first connection saved, make it active automatically.
    if not config.get("active_airlock_profile_id"):
        config["active_airlock_profile_id"] = profile["id"]
    save_config(config)

    return jsonify({"success": True, "id": profile["id"], "key_preview": key_preview(api_key)})


@app.route("/airlock/profiles/update", methods=["POST"])
def airlock_profiles_update():
    data_in = request.get_json(silent=True) or {}
    profile_id = (data_in.get("id") or "").strip()

    config = load_config()
    profiles = config.get("airlock_profiles") or []
    profile = next((p for p in profiles if p.get("id") == profile_id), None)
    if not profile:
        return jsonify({"error": "Connection not found."}), 404

    if "label" in data_in:
        new_label = (data_in.get("label") or "").strip()
        if not new_label:
            return jsonify({"error": "Label cannot be empty."}), 400
        if any(
            p.get("id") != profile_id and p.get("label", "").strip().lower() == new_label.lower() for p in profiles
        ):
            return jsonify({"error": f"A connection named '{new_label}' already exists."}), 400
        profile["label"] = new_label

    if "tenant" in data_in:
        new_tenant = (data_in.get("tenant") or "").strip()
        if not new_tenant:
            return jsonify({"error": "Tenant cannot be empty."}), 400
        profile["tenant"] = new_tenant

    if "port" in data_in:
        new_port = (data_in.get("port") or "").strip()
        if new_port and not (new_port.isdigit() and 1 <= int(new_port) <= 65535):
            return jsonify({"error": "Port must be a whole number between 1 and 65535."}), 400
        profile["port"] = new_port

    new_key = (data_in.get("api_key") or "").strip()
    if new_key:
        profile["api_key"] = new_key

    save_config(config)
    return jsonify({"success": True, "key_preview": key_preview(profile.get("api_key", ""))})


@app.route("/airlock/profiles/activate", methods=["POST"])
def airlock_profiles_activate():
    data_in = request.get_json(silent=True) or {}
    profile_id = (data_in.get("id") or "").strip()

    config = load_config()
    profiles = config.get("airlock_profiles") or []
    if not any(p.get("id") == profile_id for p in profiles):
        return jsonify({"error": "Connection not found."}), 404

    config["active_airlock_profile_id"] = profile_id
    save_config(config)
    return jsonify({"success": True})


@app.route("/airlock/profiles/delete", methods=["POST"])
def airlock_profiles_delete():
    data_in = request.get_json(silent=True) or {}
    profile_id = (data_in.get("id") or "").strip()

    config = load_config()
    profiles = config.get("airlock_profiles") or []
    remaining = [p for p in profiles if p.get("id") != profile_id]
    if len(remaining) == len(profiles):
        return jsonify({"error": "Connection not found."}), 404

    config["airlock_profiles"] = remaining
    if config.get("active_airlock_profile_id") == profile_id:
        config["active_airlock_profile_id"] = remaining[0]["id"] if remaining else None
    save_config(config)
    return jsonify({"success": True, "active_id": config.get("active_airlock_profile_id")})


@app.route("/airlock/profiles/reveal", methods=["POST"])
def airlock_profiles_reveal():
    # Mirrors /config/reveal's pattern, but scoped to one saved
    # connection's key rather than a single flat config field.
    data_in = request.get_json(silent=True) or {}
    profile_id = (data_in.get("id") or "").strip()

    config = load_config()
    profiles = config.get("airlock_profiles") or []
    profile = next((p for p in profiles if p.get("id") == profile_id), None)
    if not profile:
        return jsonify({"error": "Connection not found."}), 404

    key = profile.get("api_key", "")
    if not key:
        return jsonify({"error": "No key saved for this connection."}), 404

    return jsonify({"value": key})


# --- Custom Widgets: Partner Engagement Report (Cloud config) ---

@app.route("/cloud/config", methods=["GET"])
def cloud_config_get():
    cfg = get_cloud_admin_config()
    api_key = cfg.get("api_key", "")
    return jsonify(
        {
            "base_domain": cfg.get("base_domain", ""),
            "has_key": bool(api_key),
            "key_preview": key_preview(api_key),
        }
    )


@app.route("/cloud/config", methods=["POST"])
def cloud_config_set():
    data_in = request.get_json(silent=True) or {}
    base_domain = data_in.get("base_domain")
    api_key = data_in.get("api_key")

    cloud_cfg = load_cloud_config()
    cfg = cloud_cfg.get("cloud_admin") or {}

    if base_domain is not None:
        cfg["base_domain"] = base_domain.strip()
    if api_key is not None and api_key.strip():
        cfg["api_key"] = api_key.strip()

    cloud_cfg["cloud_admin"] = cfg
    save_cloud_config(cloud_cfg)

    return jsonify({"success": True, "base_domain": cfg.get("base_domain", ""), "key_preview": key_preview(cfg.get("api_key", ""))})


@app.route("/cloud/config/reveal", methods=["POST"])
def cloud_config_reveal():
    cfg = get_cloud_admin_config()
    key = cfg.get("api_key", "")
    if not key:
        return jsonify({"error": "No key saved."}), 404
    return jsonify({"value": key})


@app.route("/cloud/tenants", methods=["GET"])
def cloud_tenants_list():
    cloud_cfg = load_cloud_config()
    return jsonify({"tenants": cloud_cfg.get("cloud_tenants") or []})


@app.route("/cloud/tenants", methods=["POST"])
def cloud_tenants_add():
    data_in = request.get_json(silent=True) or {}
    label = (data_in.get("label") or "").strip()
    tenant_id = (data_in.get("tenant_id") or "").strip()
    directory_id = (data_in.get("directory_id") or "").strip()

    if not label:
        return jsonify({"error": "A label is required."}), 400
    if not tenant_id:
        return jsonify({"error": "A Tenant ID is required."}), 400
    if not directory_id:
        return jsonify({"error": "A Directory ID is required."}), 400

    cloud_cfg = load_cloud_config()
    tenants = cloud_cfg.get("cloud_tenants") or []

    if any(t.get("label", "").strip().lower() == label.lower() for t in tenants):
        return jsonify({"error": f"A tenant named '{label}' already exists."}), 400

    entry = {"id": str(uuid.uuid4()), "label": label, "tenant_id": tenant_id, "directory_id": directory_id}
    tenants.append(entry)
    cloud_cfg["cloud_tenants"] = tenants
    log.info("Adding NFR tenant '%s' (%s total after add)", label, len(tenants))
    save_cloud_config(cloud_cfg)

    return jsonify({"success": True, "id": entry["id"]})


@app.route("/cloud/tenants/update", methods=["POST"])
def cloud_tenants_update():
    data_in = request.get_json(silent=True) or {}
    entry_id = (data_in.get("id") or "").strip()

    cloud_cfg = load_cloud_config()
    tenants = cloud_cfg.get("cloud_tenants") or []
    entry = next((t for t in tenants if t.get("id") == entry_id), None)
    if not entry:
        return jsonify({"error": "Tenant not found."}), 404

    if "label" in data_in:
        new_label = (data_in.get("label") or "").strip()
        if not new_label:
            return jsonify({"error": "Label cannot be empty."}), 400
        if any(t.get("id") != entry_id and t.get("label", "").strip().lower() == new_label.lower() for t in tenants):
            return jsonify({"error": f"A tenant named '{new_label}' already exists."}), 400
        entry["label"] = new_label
    if "tenant_id" in data_in:
        new_tid = (data_in.get("tenant_id") or "").strip()
        if not new_tid:
            return jsonify({"error": "Tenant ID cannot be empty."}), 400
        entry["tenant_id"] = new_tid
    if "directory_id" in data_in:
        new_did = (data_in.get("directory_id") or "").strip()
        if not new_did:
            return jsonify({"error": "Directory ID cannot be empty."}), 400
        entry["directory_id"] = new_did

    save_cloud_config(cloud_cfg)
    return jsonify({"success": True})


@app.route("/cloud/tenants/delete", methods=["POST"])
def cloud_tenants_delete():
    data_in = request.get_json(silent=True) or {}
    entry_id = (data_in.get("id") or "").strip()

    cloud_cfg = load_cloud_config()
    tenants = cloud_cfg.get("cloud_tenants") or []
    remaining = [t for t in tenants if t.get("id") != entry_id]
    if len(remaining) == len(tenants):
        return jsonify({"error": "Tenant not found."}), 404

    cloud_cfg["cloud_tenants"] = remaining
    log.info("Removing NFR tenant %s (%s total after remove)", entry_id, len(remaining))
    save_cloud_config(cloud_cfg)
    return jsonify({"success": True})


# --- Custom Widgets: Timed Audit Mode ---

@app.route("/airlock/groups", methods=["GET"])
def airlock_groups():
    data, err = airlock_request("/v1/group")
    if err:
        return jsonify({"error": err, "groups": []}), 400

    # Be defensive about the exact response shape — normally it's
    # {"groups": [...]}, but fall back to finding any list in the
    # response, or treating the response itself as the list.
    if isinstance(data, dict):
        groups = data.get("groups")
        if groups is None:
            groups = next((v for v in data.values() if isinstance(v, list)), [])
    elif isinstance(data, list):
        groups = data
    else:
        groups = []

    result = []
    for g in groups:
        if not isinstance(g, dict):
            continue
        groupid = g.get("groupid") or g.get("id") or g.get("group_id")
        if not groupid:
            continue
        name = (
            g.get("name") or g.get("groupname") or g.get("group_name") or g.get("title")
            or f"(unnamed — {groupid[:8]})"
        )
        result.append({"groupid": groupid, "name": name})
    result.sort(key=lambda g: g["name"].lower())

    response = {"groups": result}
    if not result and groups:
        # Airlock returned data but nothing matched the fields we expect —
        # echo a raw sample so the mismatch is visible right in the widget
        # rather than requiring a look at the server's terminal log.
        first = groups[0]
        response["raw_sample"] = first if isinstance(first, dict) else str(first)
        log.warning("Airlock /v1/group returned data but no groups parsed. Raw sample: %s", first)

    return jsonify(response)


@app.route("/airlock/agents", methods=["POST"])
def airlock_agents():
    data_in = request.get_json(silent=True) or {}
    groupid = data_in.get("groupid", "").strip()
    if not groupid:
        return jsonify({"error": "No groupid provided.", "agents": []}), 400

    data, err = airlock_request("/v1/agent/find", {"groupid": groupid})
    if err:
        return jsonify({"error": err, "agents": []}), 400

    if isinstance(data, dict):
        agents = data.get("agents")
        if agents is None:
            agents = next((v for v in data.values() if isinstance(v, list)), [])
    elif isinstance(data, list):
        agents = data
    else:
        agents = []

    result = []
    for a in agents:
        if not isinstance(a, dict):
            continue
        agentid = a.get("agentid") or a.get("id")
        if not agentid:
            continue
        # Defensive filter: only include agents actually in the requested
        # group, in case the API doesn't filter server-side on this field.
        agent_groupid = a.get("groupid")
        if agent_groupid and agent_groupid != groupid:
            continue
        result.append(
            {
                "agentid": agentid,
                "hostname": a.get("hostname") or agentid,
                "username": a.get("username", ""),
                "os": a.get("os", ""),
            }
        )
    result.sort(key=lambda a: a["hostname"].lower())

    response = {"agents": result}
    if not result and agents:
        first = agents[0]
        response["raw_sample"] = first if isinstance(first, dict) else str(first)
        log.warning("Airlock /v1/agent/find returned data but no agents parsed. Raw sample: %s", first)

    return jsonify(response)


def process_due_audit_sessions():
    """Reverts any Timed Audit Mode session whose timer has expired.
    Also retries sessions stuck in 'revert_failed' from a prior attempt —
    a failed revert should keep being retried, not silently give up,
    since that would leave endpoints in audit mode indefinitely."""
    sessions = load_audit_sessions()
    now = datetime.now(timezone.utc)
    changed = False

    for s in sessions:
        if s.get("status") not in ("active", "revert_failed"):
            continue
        try:
            expires_at = datetime.fromisoformat(s["expires_at"])
        except (KeyError, ValueError, TypeError):
            continue
        if expires_at > now:
            continue

        ok, err = airlock_move_agents(s["source_groupid"], s["agent_ids"])
        changed = True
        if ok:
            s["status"] = "reverted"
            s["reverted_at"] = now.isoformat()
            s["revert_error"] = None
        else:
            s["status"] = "revert_failed"
            s["revert_error"] = err

    if changed:
        save_audit_sessions(sessions)


def audit_scheduler_loop():
    while True:
        try:
            process_due_audit_sessions()
        except Exception as e:  # noqa: broad except — this loop must never die
            print(f"[audit scheduler] error: {e}")
        time.sleep(20)


@app.route("/audit/sessions", methods=["GET"])
def audit_sessions_list():
    return jsonify({"sessions": load_audit_sessions()})


@app.route("/audit/start", methods=["POST"])
def audit_start():
    data_in = request.get_json(silent=True) or {}
    source_groupid = (data_in.get("source_groupid") or "").strip()
    source_group_name = (data_in.get("source_group_name") or "").strip() or source_groupid
    dest_groupid = (data_in.get("dest_groupid") or "").strip()
    dest_group_name = (data_in.get("dest_group_name") or "").strip() or dest_groupid
    agent_ids = data_in.get("agent_ids") or []
    agent_labels = data_in.get("agent_labels") or []
    duration_seconds = data_in.get("duration_seconds")

    if not source_groupid or not dest_groupid:
        return jsonify({"error": "Source and destination groups are required."}), 400
    if source_groupid == dest_groupid:
        return jsonify({"error": "Source and destination groups must be different."}), 400
    if not isinstance(agent_ids, list) or not agent_ids or not all(isinstance(a, str) and a for a in agent_ids):
        return jsonify({"error": "Select at least one agent to move."}), 400
    if not isinstance(duration_seconds, (int, float)) or duration_seconds <= 0:
        return jsonify({"error": "duration_seconds must be a positive number."}), 400

    ok, err = airlock_move_agents(dest_groupid, agent_ids)
    if not ok:
        return jsonify({"error": f"Move to '{dest_group_name}' failed: {err}"}), 502

    now = datetime.now(timezone.utc)
    session = {
        "id": str(uuid.uuid4()),
        "source_groupid": source_groupid,
        "source_group_name": source_group_name,
        "dest_groupid": dest_groupid,
        "dest_group_name": dest_group_name,
        "agent_ids": agent_ids,
        "agent_labels": agent_labels if len(agent_labels) == len(agent_ids) else agent_ids,
        "created_at": now.isoformat(),
        "expires_at": (now + timedelta(seconds=duration_seconds)).isoformat(),
        "status": "active",
        "revert_error": None,
        "reverted_at": None,
    }

    sessions = load_audit_sessions()
    sessions.append(session)
    save_audit_sessions(sessions)

    return jsonify({"success": True, "session": session})


@app.route("/audit/cancel", methods=["POST"])
def audit_cancel():
    data_in = request.get_json(silent=True) or {}
    session_id = (data_in.get("id") or "").strip()
    if not session_id:
        return jsonify({"error": "No session id provided."}), 400

    sessions = load_audit_sessions()
    session = next((s for s in sessions if s.get("id") == session_id), None)
    if not session:
        return jsonify({"error": "Session not found."}), 404
    if session.get("status") not in ("active", "revert_failed"):
        return jsonify({"error": f"Session is already '{session.get('status')}'."}), 400

    ok, err = airlock_move_agents(session["source_groupid"], session["agent_ids"])
    if not ok:
        session["status"] = "revert_failed"
        session["revert_error"] = err
        save_audit_sessions(sessions)
        return jsonify({"error": f"Revert failed: {err}"}), 502

    session["status"] = "reverted"
    session["reverted_at"] = datetime.now(timezone.utc).isoformat()
    session["revert_error"] = None
    save_audit_sessions(sessions)

    return jsonify({"success": True, "session": session})


# --- Custom Widgets: ISO 27001 Compliance Assessment ---

@app.route("/iso/mapping", methods=["GET"])
def iso_mapping_get():
    mapping, err = load_iso_mapping()
    if err:
        return jsonify({"error": err, "controls": []}), 500
    return jsonify(mapping)


CATEGORY_FULLY = "fully_compliant"
CATEGORY_PARTIAL = "partially_compliant"
CATEGORY_NON = "non_compliant"
CATEGORY_UNKNOWN = "not_assessed"

CATEGORY_LABELS = {
    CATEGORY_FULLY: "Fully Compliant",
    CATEGORY_PARTIAL: "Partially Compliant",
    CATEGORY_NON: "Non-Compliant",
    CATEGORY_UNKNOWN: "Not Assessed",
}
CATEGORY_COLORS = {
    CATEGORY_FULLY: "#1e7e34",
    CATEGORY_PARTIAL: "#b8860b",
    CATEGORY_NON: "#c53929",
    CATEGORY_UNKNOWN: "#6b6b70",
}


def categorize_group(controls):
    """Rolls up a group's per-control results into one overall bucket.
    Only controls with a definitive automated result (meets/partial/
    unmet) count toward this — 'manual' controls without a confirmed
    rule yet are excluded so an unscored control can't silently drag a
    group into 'non-compliant'."""
    scored = [c for c in controls if c.get("status") in (STATUS_MEETS, STATUS_PARTIAL, STATUS_UNMET)]
    if not scored:
        return CATEGORY_UNKNOWN
    if all(c["status"] == STATUS_MEETS for c in scored):
        return CATEGORY_FULLY
    if all(c["status"] == STATUS_UNMET for c in scored):
        return CATEGORY_NON
    return CATEGORY_PARTIAL


def get_group_agent_count(groupid):
    """Returns (count, error)."""
    data, err = airlock_request("/v1/agent/find", {"groupid": groupid})
    if err:
        return None, err
    # Some Airlock responses use an explicit null (not a missing key or
    # empty list) for "agents" when a group has zero agents — .get()'s
    # default only covers a missing key, so 'or []' is needed to also
    # catch the null case.
    agents = (data.get("agents") if isinstance(data, dict) else None) or []
    return len(agents), None


def build_svg_donut_chart(segments, size=220, stroke_width=34):
    """segments: list of (label, pct, color) tuples with pct in 0-100.
    Returns a self-contained SVG donut chart string — no external
    dependencies or network access needed, so it renders identically in
    the live dashboard and in a downloaded, offline HTML report."""
    segments = [s for s in segments if s[1] > 0]
    radius = (size - stroke_width) / 2
    cx = cy = size / 2

    if not segments:
        return (
            f'<svg width="{size}" height="{size}" viewBox="0 0 {size} {size}" xmlns="http://www.w3.org/2000/svg">'
            f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="#eceef2" stroke-width="{stroke_width}"/>'
            f"</svg>"
        )

    circumference = 2 * math.pi * radius
    parts = [f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="#eceef2" stroke-width="{stroke_width}"/>']
    offset = 0.0
    for _, pct, color in segments:
        dash = circumference * (pct / 100.0)
        gap = circumference - dash
        parts.append(
            f'<circle cx="{cx}" cy="{cy}" r="{radius}" fill="none" stroke="{color}" '
            f'stroke-width="{stroke_width}" stroke-dasharray="{dash:.2f} {gap:.2f}" '
            f'stroke-dashoffset="{-offset:.2f}" transform="rotate(-90 {cx} {cy})"/>'
        )
        offset += dash

    return f'<svg width="{size}" height="{size}" viewBox="0 0 {size} {size}" xmlns="http://www.w3.org/2000/svg">{"".join(parts)}</svg>'


def build_compliance_summary(group_results):
    """Rolls per-group categories up into endpoint-weighted percentages
    across the whole fleet, plus a ready-to-embed SVG chart."""
    totals = {CATEGORY_FULLY: 0, CATEGORY_PARTIAL: 0, CATEGORY_NON: 0, CATEGORY_UNKNOWN: 0}
    missing_agent_count = 0

    for g in group_results:
        count = g.get("agent_count")
        if count is None:
            missing_agent_count += 1
            continue
        category = g.get("category", CATEGORY_UNKNOWN)
        totals[category] = totals.get(category, 0) + count

    total_agents = sum(totals.values())

    segments = []
    for cat in (CATEGORY_FULLY, CATEGORY_PARTIAL, CATEGORY_NON, CATEGORY_UNKNOWN):
        count = totals[cat]
        pct = round((count / total_agents * 100), 1) if total_agents else 0.0
        segments.append(
            {
                "category": cat,
                "label": CATEGORY_LABELS[cat],
                "color": CATEGORY_COLORS[cat],
                "agent_count": count,
                "percent": pct,
            }
        )

    chart_svg = build_svg_donut_chart([(s["label"], s["percent"], s["color"]) for s in segments])

    return {
        "total_agents": total_agents,
        "groups_missing_agent_count": missing_agent_count,
        "segments": segments,
        "chart_svg": chart_svg,
    }


def run_iso_assessment():
    """Pulls every Airlock policy group, evaluates each against the ISO
    mapping, and returns (result_dict, error). error is a plain message
    if the group listing itself failed; per-group policy-fetch failures
    are instead recorded per-group in the result so one bad group
    doesn't block the whole assessment."""
    mapping, err = load_iso_mapping()
    if err:
        return None, err
    if not mapping.get("controls"):
        return None, f"{os.path.basename(ISO_MAPPING_PATH)} was found and is valid JSON, but has no controls defined."

    data, err = airlock_request("/v1/group")
    if err:
        return None, err

    groups = (data.get("groups") if isinstance(data, dict) else None) or []
    groups = [g for g in groups if isinstance(g, dict) and g.get("groupid")]

    group_results = []
    for g in groups:
        groupid = g["groupid"]
        name = g.get("name") or f"(unnamed — {groupid[:8]})"

        policy_data, policy_err = airlock_request("/v1/group/policies", {"groupid": groupid})
        if policy_err:
            group_results.append(
                {
                    "groupid": groupid,
                    "name": name,
                    "error": policy_err,
                    "controls": [],
                    "category": CATEGORY_UNKNOWN,
                    "agent_count": None,
                    "agent_count_error": None,
                }
            )
            continue

        controls = evaluate_group_against_mapping(policy_data or {}, mapping)
        category = categorize_group(controls)
        agent_count, agent_count_error = get_group_agent_count(groupid)

        group_results.append(
            {
                "groupid": groupid,
                "name": name,
                "error": None,
                "controls": controls,
                "category": category,
                "agent_count": agent_count,
                "agent_count_error": agent_count_error,
            }
        )

    summary = build_compliance_summary(group_results)

    return {"controls": mapping["controls"], "groups": group_results, "summary": summary}, None


@app.route("/iso/assessment", methods=["GET"])
def iso_assessment():
    result, err = run_iso_assessment()
    if err:
        return jsonify({"error": err}), 400
    return jsonify(result)


def build_iso_report_html(assessment):
    """Builds a self-contained HTML report summarizing the ISO
    assessment. Returns the HTML as a string — light-background and
    print-friendly, since this is meant to be handed to a client (and
    can be printed to PDF straight from the browser if needed)."""
    status_colors = {
        STATUS_MEETS: ("#1e7e34", "#e3f5ea"),
        STATUS_PARTIAL: ("#8a6d1b", "#fdf3d9"),
        STATUS_UNMET: ("#c53929", "#fbe6e3"),
        STATUS_UNKNOWN: ("#6b6b70", "#eceef2"),
        STATUS_ERROR: ("#6b6b70", "#eceef2"),
    }
    status_labels = {
        STATUS_MEETS: "Meets",
        STATUS_PARTIAL: "Partial",
        STATUS_UNMET: "Not met",
        STATUS_UNKNOWN: "Unknown",
        STATUS_ERROR: "Error",
    }

    def esc(value):
        return html.escape(str(value if value is not None else ""))

    def badge(status):
        color, bg = status_colors.get(status, status_colors[STATUS_UNKNOWN])
        label = status_labels.get(status, status)
        return f'<span class="badge" style="color:{color};background:{bg};">{esc(label)}</span>'

    controls = assessment.get("controls", [])
    groups = assessment.get("groups", [])
    summary = assessment.get("summary", {})

    category_colors = CATEGORY_COLORS
    category_labels = CATEGORY_LABELS

    def category_badge(category):
        color = category_colors.get(category, "#6b6b70")
        label = category_labels.get(category, category)
        return f'<span class="badge" style="color:#fff;background:{color};">{esc(label)}</span>'

    controls_html = "".join(
        f'<div class="control-summary">'
        f'<div class="control-summary-title">{esc(c.get("title"))} — {esc(c.get("name"))}</div>'
        f'<div class="control-summary-desc">{esc(c.get("description"))}</div>'
        f"</div>"
        for c in controls
    )

    group_blocks = []
    for g in groups:
        if g.get("error"):
            group_blocks.append(
                f'<div class="group-card"><div class="group-name">{esc(g.get("name"))}</div>'
                f'<div class="group-error">Couldn\'t assess this group: {esc(g.get("error"))}</div></div>'
            )
            continue

        agent_count = g.get("agent_count")
        agent_note = f"{agent_count} endpoint(s)" if agent_count is not None else "endpoint count unavailable"

        rows = "".join(
            f"<tr><td><strong>{esc(cr.get('title'))}</strong> — {esc(cr.get('name'))}"
            f'<div class="detail">{esc(cr.get("detail"))}</div></td>'
            f"<td>{badge(cr.get('status'))}</td></tr>"
            for cr in g.get("controls", [])
        )
        group_blocks.append(
            f'<div class="group-card">'
            f'<div class="group-name">{esc(g.get("name"))} {category_badge(g.get("category"))}</div>'
            f'<div class="group-meta">{esc(agent_note)}</div>'
            f"<table>{rows}</table></div>"
        )
    groups_html = "".join(group_blocks)

    chart_svg = summary.get("chart_svg", "")
    legend_html = "".join(
        f'<div class="legend-row"><span class="legend-swatch" style="background:{s["color"]};"></span>'
        f'{esc(s["label"])}: <strong>{s["percent"]}%</strong> ({s["agent_count"]} endpoint{"s" if s["agent_count"] != 1 else ""})</div>'
        for s in summary.get("segments", [])
    )
    missing = summary.get("groups_missing_agent_count", 0)
    missing_note = f" {missing} group(s) excluded — endpoint count unavailable." if missing else ""
    chart_section = f"""
  <h2>Endpoint compliance breakdown</h2>
  <div class="chart-row">
    <div class="chart-svg">{chart_svg}</div>
    <div class="chart-legend">
      {legend_html}
      <div class="chart-meta">Based on {summary.get("total_agents", 0)} managed endpoint(s) across {len(groups)} group(s).{esc(missing_note)}</div>
    </div>
  </div>
"""

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>ISO 27001 Alignment — Airlock Digital Policy Assessment</title>
<style>
  body {{
    font-family: -apple-system, "Segoe UI", Helvetica, Arial, sans-serif;
    background: #ffffff;
    color: #1e1e1e;
    max-width: 900px;
    margin: 40px auto;
    padding: 0 20px 60px;
    line-height: 1.5;
  }}
  h1 {{ font-size: 22px; margin-bottom: 4px; }}
  h2 {{ font-size: 17px; border-bottom: 1px solid #d0d0d5; padding-bottom: 6px; margin-top: 32px; }}
  .meta {{ color: #6b6b70; font-size: 13px; margin-bottom: 20px; }}
  .disclaimer {{
    background: #f4f4f5;
    border-left: 4px solid #5b9dff;
    padding: 12px 16px;
    font-size: 13px;
    color: #444;
    margin-bottom: 28px;
  }}
  .control-summary {{ margin-bottom: 14px; }}
  .control-summary-title {{ font-weight: 600; font-size: 14px; }}
  .control-summary-desc {{ color: #555; font-size: 13px; margin-top: 2px; }}
  .group-card {{
    border: 1px solid #d0d0d5;
    border-radius: 6px;
    padding: 16px 18px;
    margin-bottom: 16px;
    page-break-inside: avoid;
  }}
  .group-name {{ font-weight: 600; font-size: 15px; margin-bottom: 10px; }}
  .group-error {{ color: #c53929; font-size: 13px; }}
  .group-meta {{ color: #6b6b70; font-size: 12px; margin-bottom: 10px; }}
  .chart-row {{
    display: flex;
    align-items: center;
    gap: 28px;
    flex-wrap: wrap;
    margin-bottom: 8px;
  }}
  .chart-legend {{ flex: 1; min-width: 220px; }}
  .legend-row {{ display: flex; align-items: center; gap: 8px; font-size: 13px; margin-bottom: 6px; }}
  .legend-swatch {{ width: 12px; height: 12px; border-radius: 3px; flex-shrink: 0; }}
  .chart-meta {{ color: #6b6b70; font-size: 12px; margin-top: 8px; }}
  table {{ width: 100%; border-collapse: collapse; }}
  td {{ padding: 8px 4px; border-top: 1px solid #eceef2; font-size: 13px; vertical-align: top; }}
  tr:first-child td {{ border-top: none; }}
  .detail {{ color: #6b6b70; font-size: 12px; margin-top: 2px; }}
  .badge {{
    display: inline-block;
    font-size: 11px;
    font-weight: 600;
    padding: 3px 10px;
    border-radius: 10px;
    text-transform: uppercase;
    letter-spacing: 0.03em;
    white-space: nowrap;
  }}
  @media print {{
    body {{ margin: 0 auto; }}
    .group-card {{ break-inside: avoid; }}
  }}
</style>
</head>
<body>
  <h1>ISO 27001 Alignment — Airlock Digital Policy Assessment</h1>
  <div class="meta">Generated: {esc(generated_at)}</div>
  <div class="disclaimer">
    This report shows how Airlock Digital's policy configuration provides evidence toward
    specific ISO/IEC 27001:2022 Annex A controls. It covers only the subset of controls that
    an endpoint application-control product can speak to directly — it is not a substitute for
    a full ISMS audit, and does not by itself constitute or guarantee ISO 27001 certification.
  </div>
  {chart_section}
  <h2>Controls assessed</h2>
  {controls_html}

  <h2>Results by policy group</h2>
  {groups_html}
</body>
</html>
"""


@app.route("/iso/report", methods=["GET"])
def iso_report():
    result, err = run_iso_assessment()
    if err:
        return jsonify({"error": err}), 400

    try:
        html_report = build_iso_report_html(result)
    except Exception as e:
        return jsonify({"error": f"Couldn't generate report: {e}"}), 500

    filename = f"ISO27001_Airlock_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.html"
    response = app.response_class(html_report, mimetype="text/html")
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def build_cloud_partner_report_html(report_data):
    """Builds a self-contained HTML report summarizing partner tenant
    engagement/utilization across the Cloud multi-tenant environment.
    Opens directly in a browser tab rather than downloading, so this
    stays inline (no Content-Disposition: attachment)."""

    def esc(value):
        return html.escape(str(value if value is not None else ""))

    def fmt_date(iso_str):
        if not iso_str:
            return "Never"
        try:
            dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
            return dt.strftime("%Y-%m-%d %H:%M UTC")
        except (ValueError, TypeError):
            return str(iso_str)

    tenants = report_data.get("tenants", [])
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    summary_rows = []
    detail_sections = []

    for t in tenants:
        label = t.get("label", "")

        total_agents = t.get("client_count", 0)
        allocated = t.get("license_allocated")
        has_license_data = allocated is not None and not t.get("license_error")
        utilization = round((total_agents / allocated * 100), 1) if has_license_data and allocated else None
        users = t.get("users", [])
        last_logins = [u.get("last_accessed") for u in users if u.get("last_accessed")]
        most_recent = max(last_logins) if last_logins else ""

        agents_cell = str(total_agents) if not t.get("agents_error") else '<span class="error-cell">Error</span>'
        license_cell = (
            f"{total_agents}/{allocated}" if has_license_data else '<span class="error-cell">Unavailable</span>'
        )
        utilization_cell = f"{utilization}%" if utilization is not None else "—"
        users_cell = str(len(users)) if not t.get("users_error") else '<span class="error-cell">Error</span>'
        login_cell = esc(fmt_date(most_recent)) if not t.get("users_error") else "—"

        summary_rows.append(
            f"<tr>"
            f"<td>{esc(label)}</td>"
            f"<td>{agents_cell}</td>"
            f"<td>{t.get('enforce_count', 0)}</td>"
            f"<td>{t.get('audit_count', 0)}</td>"
            f"<td>{license_cell}</td>"
            f"<td>{utilization_cell}</td>"
            f"<td>{users_cell}</td>"
            f"<td>{login_cell}</td>"
            f"</tr>"
        )

        users_sorted = sorted(users, key=lambda u: u.get("last_accessed") or "", reverse=True)
        user_rows = "".join(
            f"<tr><td>{esc(u.get('full_name'))}</td><td>{esc(u.get('email'))}</td>"
            f"<td>{esc(fmt_date(u.get('last_accessed')))}</td></tr>"
            for u in users_sorted
        )

        error_notes = []
        if t.get("users_error"):
            error_notes.append(f"Users: {t['users_error']}")
        if t.get("agents_error"):
            error_notes.append(f"Agent counts: {t['agents_error']}")
        if t.get("license_error"):
            error_notes.append(f"License data: {t['license_error']}")
        if t.get("policy_errors"):
            error_notes.append("Policy breakdown: " + "; ".join(t["policy_errors"]))
        error_block = (
            "".join(f'<p class="error-cell">{esc(n)}</p>' for n in error_notes) if error_notes else ""
        )

        license_summary = f"{total_agents}/{allocated} licenses used ({utilization}%)" if has_license_data else "license data unavailable"

        detail_sections.append(
            f"""<details>
  <summary>{esc(label)} &mdash; {agents_cell if not t.get('agents_error') else '?'} agent(s), {len(users) if not t.get('users_error') else '?'} user(s)</summary>
  <div class="detail-body">
    <div class="detail-stats">
      <span><strong>{t.get('enforce_count', 0)}</strong> enforcing</span>
      <span><strong>{t.get('audit_count', 0)}</strong> audit-only</span>
      <span><strong>{t.get('unmanaged_count', 0)}</strong> unmanaged</span>
      <span>{esc(license_summary)}</span>
    </div>
    {error_block}
    <table>
      <thead><tr><th>User</th><th>Email</th><th>Last login</th></tr></thead>
      <tbody>{user_rows}</tbody>
    </table>
  </div>
</details>"""
        )

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Partner Engagement &amp; Utilization Report</title>
<style>
  /* Same palette as the Partner Consulting Toolkit console itself,
     so this report feels like it came out of the same app. */
  :root {{
    --bg: #12173c;
    --panel: #1b2358;
    --border: #333e82;
    --text: #e8eaf6;
    --accent: #5b9dff;
    --muted: #8a91c4;
  }}
  body {{
    font-family: -apple-system, "Segoe UI", Helvetica, Arial, sans-serif;
    background: var(--bg);
    color: var(--text);
    max-width: 960px;
    margin: 40px auto;
    padding: 0 20px 60px;
    line-height: 1.5;
  }}
  h1 {{ font-size: 22px; margin-bottom: 4px; color: var(--text); }}
  h2 {{ font-size: 17px; border-bottom: 1px solid var(--border); padding-bottom: 6px; margin-top: 32px; color: var(--text); }}
  .meta {{ color: var(--muted); font-size: 13px; margin-bottom: 24px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 12px; background: var(--panel); border-radius: 6px; overflow: hidden; }}
  th, td {{ padding: 8px 10px; border-top: 1px solid var(--border); font-size: 13px; text-align: left; vertical-align: top; }}
  thead th {{ border-top: none; border-bottom: 2px solid var(--accent); color: var(--text); background: rgba(91, 157, 255, 0.1); }}
  tr:first-child td {{ border-top: none; }}
  .error-cell {{ color: #ff8a8a; }}
  details {{
    border: 1px solid var(--border);
    border-radius: 6px;
    padding: 12px 16px;
    margin-bottom: 12px;
    background: var(--panel);
  }}
  summary {{ cursor: pointer; font-weight: 600; font-size: 14px; color: var(--text); }}
  summary::marker {{ color: var(--accent); }}
  .detail-body {{ margin-top: 12px; }}
  .detail-stats {{ display: flex; gap: 20px; flex-wrap: wrap; font-size: 13px; color: var(--muted); margin-bottom: 8px; }}
  .detail-stats strong {{ color: var(--accent); }}
  @media print {{
    :root {{
      --bg: #ffffff;
      --panel: #ffffff;
      --border: #d0d0d5;
      --text: #1e1e1e;
      --accent: #2f5fae;
      --muted: #6b6b70;
    }}
    body {{ margin: 0 auto; }}
    table {{ background: none; }}
    thead th {{ background: none; }}
    details {{ break-inside: avoid; background: none; }}
    details:not([open]) summary ~ * {{ display: block !important; }}
  }}
</style>
</head>
<body>
  <h1>Partner Engagement &amp; Utilization Report</h1>
  <div class="meta">Generated: {esc(generated_at)}</div>

  <h2>Summary</h2>
  <table>
    <thead>
      <tr><th>Tenant</th><th>Total agents</th><th>Enforce</th><th>Audit</th><th>Licenses used/allocated</th><th>Utilization</th><th>Users</th><th>Most recent login</th></tr>
    </thead>
    <tbody>
      {"".join(summary_rows)}
    </tbody>
  </table>

  <h2>Tenant detail</h2>
  {"".join(detail_sections)}
</body>
</html>
"""


def _cloud_report_fmt_date(iso_str):
    """Shared date formatting for the xlsx/xml exports — same behavior
    as the inline fmt_date() in build_cloud_partner_report_html."""
    if not iso_str:
        return "Never"
    try:
        dt = datetime.fromisoformat(iso_str.replace("Z", "+00:00"))
        return dt.strftime("%Y-%m-%d %H:%M UTC")
    except (ValueError, TypeError):
        return str(iso_str)


def build_cloud_partner_report_workbook(report_data):
    """Builds an .xlsx workbook mirroring the HTML report: a Summary
    sheet across all tenants, plus one sheet per tenant with its full
    user list. Returns a BytesIO ready to hand to send_file()."""
    tenants = report_data.get("tenants", [])
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B2358", end_color="1B2358", fill_type="solid")

    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"

    summary_ws.append(["Partner Engagement & Utilization Report"])
    summary_ws["A1"].font = Font(bold=True, size=14)
    summary_ws.append([f"Generated: {generated_at}"])
    summary_ws.append([])

    headers = [
        "Tenant", "Total agents", "Enforce", "Audit", "Licenses used",
        "Licenses allocated", "Utilization %", "Users", "Most recent login", "Notes",
    ]
    header_row = summary_ws.max_row + 1
    summary_ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = summary_ws.cell(row=header_row, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for t in tenants:
        total_agents = t.get("client_count", 0)
        allocated = t.get("license_allocated")
        has_license_data = allocated is not None and not t.get("license_error")
        utilization = round((total_agents / allocated * 100), 1) if has_license_data and allocated else None
        users = t.get("users", [])
        last_logins = [u.get("last_accessed") for u in users if u.get("last_accessed")]
        most_recent = max(last_logins) if last_logins else ""

        notes = []
        if t.get("users_error"):
            notes.append(f"Users: {t['users_error']}")
        if t.get("agents_error"):
            notes.append(f"Agent counts: {t['agents_error']}")
        if t.get("license_error"):
            notes.append(f"License data: {t['license_error']}")
        if t.get("policy_errors"):
            notes.append("Policy breakdown: " + "; ".join(t["policy_errors"]))

        summary_ws.append([
            t.get("label", ""),
            "Error" if t.get("agents_error") else total_agents,
            t.get("enforce_count", 0),
            t.get("audit_count", 0),
            total_agents if has_license_data else "Unavailable",
            allocated if has_license_data else "Unavailable",
            utilization if utilization is not None else "—",
            "Error" if t.get("users_error") else len(users),
            "—" if t.get("users_error") else _cloud_report_fmt_date(most_recent),
            "; ".join(notes),
        ])

    summary_ws.column_dimensions["A"].width = 24
    for col_idx, header in enumerate(headers[1:], start=2):
        summary_ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(header) + 4)
    summary_ws.column_dimensions[get_column_letter(len(headers))].width = 50

    # One sheet per tenant with the full user list — sheet names are
    # capped at 31 chars and can't collide, so sanitize and dedupe.
    used_sheet_names = {"Summary"}
    for t in tenants:
        label = t.get("label") or "Tenant"
        base_name = re.sub(r"[\[\]:*?/\\]", "_", label)[:31] or "Tenant"
        safe_name = base_name
        suffix = 2
        while safe_name in used_sheet_names:
            trim = 31 - len(f" ({suffix})")
            safe_name = f"{base_name[:trim]} ({suffix})"
            suffix += 1
        used_sheet_names.add(safe_name)

        ws = wb.create_sheet(title=safe_name)
        ws.append([f"{label} — Users"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        user_header_row = ws.max_row + 1
        ws.append(["Full name", "Email", "Last login"])
        for col_idx in range(1, 4):
            cell = ws.cell(row=user_header_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        users_sorted = sorted(t.get("users", []), key=lambda u: u.get("last_accessed") or "", reverse=True)
        for u in users_sorted:
            ws.append([u.get("full_name", ""), u.get("email", ""), _cloud_report_fmt_date(u.get("last_accessed"))])

        if t.get("users_error"):
            ws.append([])
            ws.append([f"Users unavailable: {t['users_error']}"])

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 32
        ws.column_dimensions["C"].width = 22

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_cloud_partner_report_xml(report_data):
    """Builds an XML document mirroring the HTML/Excel reports — one
    <Tenant> element per partner tenant, with its users nested inside
    and any per-field errors called out explicitly rather than just
    omitted, same philosophy as the HTML report's error notes."""
    tenants = report_data.get("tenants", [])
    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M")

    root = ET.Element("PartnerEngagementReport", {"generated": generated_at})

    for t in tenants:
        total_agents = t.get("client_count", 0)
        allocated = t.get("license_allocated")
        has_license_data = allocated is not None and not t.get("license_error")
        utilization = round((total_agents / allocated * 100), 1) if has_license_data and allocated else None
        users = t.get("users", [])
        last_logins = [u.get("last_accessed") for u in users if u.get("last_accessed")]
        most_recent = max(last_logins) if last_logins else ""

        tenant_el = ET.SubElement(root, "Tenant", {
            "label": t.get("label", ""),
            "tenantId": t.get("tenant_id", ""),
            "directoryId": t.get("directory_id", ""),
        })

        ET.SubElement(tenant_el, "TotalAgents", {"error": "true" if t.get("agents_error") else "false"}).text = str(total_agents)
        ET.SubElement(tenant_el, "Enforce").text = str(t.get("enforce_count", 0))
        ET.SubElement(tenant_el, "Audit").text = str(t.get("audit_count", 0))

        license_el = ET.SubElement(tenant_el, "License", {"available": "true" if has_license_data else "false"})
        ET.SubElement(license_el, "Used").text = str(total_agents) if has_license_data else ""
        ET.SubElement(license_el, "Allocated").text = str(allocated) if has_license_data else ""
        ET.SubElement(license_el, "UtilizationPercent").text = "" if utilization is None else str(utilization)

        ET.SubElement(tenant_el, "MostRecentLogin").text = "" if t.get("users_error") else _cloud_report_fmt_date(most_recent)

        users_el = ET.SubElement(tenant_el, "Users", {"error": t.get("users_error") or ""})
        for u in sorted(users, key=lambda u: u.get("last_accessed") or "", reverse=True):
            user_el = ET.SubElement(users_el, "User")
            ET.SubElement(user_el, "FullName").text = u.get("full_name") or ""
            ET.SubElement(user_el, "Email").text = u.get("email") or ""
            ET.SubElement(user_el, "LastLogin").text = _cloud_report_fmt_date(u.get("last_accessed"))

        errors_el = ET.SubElement(tenant_el, "Errors")
        if t.get("agents_error"):
            ET.SubElement(errors_el, "Error", {"field": "agent_counts"}).text = t["agents_error"]
        if t.get("license_error"):
            ET.SubElement(errors_el, "Error", {"field": "license"}).text = t["license_error"]
        if t.get("users_error"):
            ET.SubElement(errors_el, "Error", {"field": "users"}).text = t["users_error"]
        for perr in t.get("policy_errors") or []:
            ET.SubElement(errors_el, "Error", {"field": "policy"}).text = perr

    raw = ET.tostring(root, encoding="unicode")
    pretty = minidom.parseString(raw).toprettyxml(indent="  ")
    # minidom's toprettyxml scatters blank lines between elements; drop
    # them so the output is clean without losing the indentation.
    return "\n".join(line for line in pretty.split("\n") if line.strip()) + "\n"


@app.route("/cloud/report/export/xlsx", methods=["POST"])
def cloud_report_export_xlsx():
    data_in = request.get_json(silent=True) or {}
    tenant_ids = data_in.get("tenant_ids")

    result, err = run_cloud_partner_report(tenant_ids=tenant_ids)
    if err:
        return jsonify({"error": err}), 400

    try:
        buf = build_cloud_partner_report_workbook(result)
    except Exception as e:
        return jsonify({"error": f"Couldn't generate Excel export: {e}"}), 500

    filename = f"Partner_Engagement_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/cloud/report/export/xml", methods=["POST"])
def cloud_report_export_xml():
    data_in = request.get_json(silent=True) or {}
    tenant_ids = data_in.get("tenant_ids")

    result, err = run_cloud_partner_report(tenant_ids=tenant_ids)
    if err:
        return jsonify({"error": err}), 400

    try:
        xml_report = build_cloud_partner_report_xml(result)
    except Exception as e:
        return jsonify({"error": f"Couldn't generate XML export: {e}"}), 500

    filename = f"Partner_Engagement_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xml"
    response = app.response_class(xml_report, mimetype="application/xml")
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@app.route("/cloud/report", methods=["POST"])
def cloud_report():
    data_in = request.get_json(silent=True) or {}
    tenant_ids = data_in.get("tenant_ids")

    result, err = run_cloud_partner_report(tenant_ids=tenant_ids)
    if err:
        return jsonify({"error": err}), 400

    try:
        html_report = build_cloud_partner_report_html(result)
    except Exception as e:
        return jsonify({"error": f"Couldn't generate report: {e}"}), 500

    # Opens directly in a browser tab, unlike the ISO report's forced
    # download — no Content-Disposition header.
    return app.response_class(html_report, mimetype="text/html")


if __name__ == "__main__":
    # Apply the saved logging preference — configure_logging(True) ran
    # at import time as a safe default, this reflects whatever the user
    # actually last chose in Settings.
    _startup_config = load_config()
    _logging_pref = _startup_config.get("logging_enabled")
    configure_logging(True if _logging_pref is None else _logging_pref.strip().lower() == "true")

    # Print exactly where this instance reads/writes its config files —
    # the fastest way to catch a "running from the wrong folder" mixup,
    # which looks identical to a data-loss bug from the outside.
    print(f"[startup] Working directory: {BASE_DIR}")
    print(f"[startup] Main config: {CONFIG_PATH}")
    print(f"[startup] NFR Tracking config: {CLOUD_CONFIG_PATH}")

    # One-time migration for anyone upgrading from a single flat
    # Airlock tenant/key into the multi-connection system.
    migrate_legacy_airlock_profile()

    # One-time migration: NFR Tracking's admin credential and tenant
    # list used to live inside the main config.json — move them into
    # their own dedicated file.
    migrate_legacy_cloud_config()

    # If GitHub sync is enabled, pull the latest scripts down before serving
    # so the API Calls tab reflects the current repo state on startup.
    startup_sync = sync_github_scripts()
    if startup_sync.get("synced"):
        print(f"[startup sync] Pulled {len(startup_sync['downloaded'])} script(s) from {startup_sync['location']}")
    elif startup_sync.get("reason") not in ("disabled", "no repo configured"):
        print(f"[startup sync] Skipped: {startup_sync.get('reason')}")

    # Catch up on any Timed Audit Mode sessions that expired while the app
    # wasn't running, then start the background thread that watches for
    # future expirations.
    process_due_audit_sessions()
    threading.Thread(target=audit_scheduler_loop, daemon=True).start()

    # host="127.0.0.1" keeps this reachable only from your own machine.
    app.run(host="127.0.0.1", port=5000, debug=False)
